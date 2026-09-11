from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

import solve_question_one as q1


def run_case(
    df: pd.DataFrame,
    *,
    scenario_group: str,
    scenario_name: str,
    eta: float,
    soc_min: float,
    soc_max: float,
    soc_initial: float,
    power_limit_kw: float,
    mip_gap: float,
    time_limit: float | None,
    solver_python: Path | None,
) -> dict[str, object]:
    result, info = q1.solve_dispatch(
        df,
        method="milp",
        objective_mode="cost",
        eta=eta,
        soc_min=soc_min,
        soc_max=soc_max,
        soc_initial=soc_initial,
        soc_terminal=soc_initial,
        power_limit_kw=power_limit_kw,
        mip_gap=mip_gap,
        time_limit=time_limit,
        solver_python=solver_python,
    )
    checks = q1.validate_solution(
        result,
        energy_limit=power_limit_kw * q1.DELTA_H,
        eta=eta,
        soc_min=soc_min,
        soc_max=soc_max,
        soc_terminal=soc_initial,
    )
    violation_keys = (
        "power_violations",
        "curtailment_violations",
        "soc_lower_bound_violations",
        "soc_upper_bound_violations",
        "simultaneous_charge_discharge_periods",
    )
    if checks["max_balance_error"] > 1e-6 or checks["max_soc_error"] > 1e-6:
        raise RuntimeError(f"{scenario_name}的平衡或SOC递推校验失败: {checks}")
    if checks["terminal_soc_error"] > 1e-6 or any(checks[key] != 0 for key in violation_keys):
        raise RuntimeError(f"{scenario_name}违反基础模型硬约束: {checks}")
    throughput = checks["total_charge_kwh"] + checks["total_discharge_kwh"]
    return {
        "scenario_group": scenario_group,
        "scenario_name": scenario_name,
        "eta": eta,
        "soc_min": soc_min,
        "soc_max": soc_max,
        "soc_initial": soc_initial,
        "power_limit_kw": power_limit_kw,
        "cost": checks["total_cost"],
        "grid_kwh": checks["total_grid_kwh"],
        "throughput_kwh": throughput,
        "equivalent_cycles": throughput / 24000.0,
        "peak_kw": float(info["peak_kw"]),
        "min_soc": checks["min_soc"],
        "max_soc": checks["max_soc"],
        "max_charge_kwh": checks["max_charge_kwh"],
        "max_discharge_kwh": checks["max_discharge_kwh"],
        "balance_error": checks["max_balance_error"],
        "soc_error": checks["max_soc_error"],
        "power_violations": checks["power_violations"],
        "curtailment_violations": checks["curtailment_violations"],
        "soc_lower_bound_violations": checks["soc_lower_bound_violations"],
        "soc_upper_bound_violations": checks["soc_upper_bound_violations"],
        "terminal_soc_error": checks["terminal_soc_error"],
        "simultaneous_periods": checks[
            "simultaneous_charge_discharge_periods"
        ],
        "simultaneous_kwh": checks["simultaneous_charge_discharge_kwh"],
    }


def build_summary_markdown(
    frame: pd.DataFrame,
    base_cost: float,
) -> str:
    lines = [
        "# C题第一问灵敏度分析",
        "",
        "所有结果均由 MILP 重新求解得到，目标函数为全天购电费用。",
        "",
        f"基准购电费：{base_cost:.6f} 元。",
        "",
    ]
    for group in frame["scenario_group"].drop_duplicates():
        subset = frame[frame["scenario_group"] == group].copy()
        subset["delta_cost"] = subset["cost"] - base_cost
        subset["cost_change_pct"] = subset["delta_cost"] / base_cost * 100.0
        lines.extend(
            [
                f"## {group}",
                "",
                "| 情景 | 购电费/元 | 费用增量/元 | 费用变化 | 循环次数 | 峰值/kW |",
                "|---|---:|---:|---:|---:|---:|",
            ]
        )
        for _, row in subset.iterrows():
            lines.append(
                f"| {row['scenario_name']} | {row['cost']:.6f} | "
                f"{row['delta_cost']:.6f} | "
                f"{row['cost_change_pct']:.6f}% | "
                f"{row['equivalent_cycles']:.6f} | "
                f"{row['peak_kw']:.6f} |"
            )
        lines.append("")

    lines.extend(
        [
            "## 主要结论",
            "",
            "1. 负载预测是影响购电费用的首要因素，费用对负载变化近似线性放大。",
            "2. PV预测偏低会显著增加购电费，PV预测偏高时会受到储能容量、功率和弃光约束限制。",
            "3. 储能效率是第二类高敏感参数，效率下降会明显提高购电费用。",
            "4. 初始和终端SOC在可行范围内对费用影响较小。",
            "5. 最大充放电功率高于5000 kW后边际收益很小。",
            "6. 扩大可用SOC区间可降低成本，但题目给出的运行边界不应为了优化结果而修改。",
        ]
    )
    return "\n".join(lines) + "\n"


def write_results_workbook(frame: pd.DataFrame, output_path: Path) -> None:
    workbook = openpyxl.Workbook()
    results_sheet = workbook.active
    results_sheet.title = "灵敏度结果"
    results_sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        results_sheet.append(list(row))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in results_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    results_sheet.freeze_panes = "A2"
    results_sheet.auto_filter.ref = results_sheet.dimensions
    for column in results_sheet.columns:
        width = min(24, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        results_sheet.column_dimensions[column[0].column_letter].width = width

    summary_sheet = workbook.create_sheet("分组汇总")
    summary_sheet.append(["参数组", "最低费用/元", "最高费用/元", "最大费用变化/%", "情景数"])
    for group, subset in frame.groupby("scenario_group", sort=False):
        summary_sheet.append(
            [
                group,
                float(subset["cost"].min()),
                float(subset["cost"].max()),
                float(subset["cost_change_pct"].abs().max()),
                int(len(subset)),
            ]
        )
    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions
    for column in summary_sheet.columns:
        summary_sheet.column_dimensions[column[0].column_letter].width = 20

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="C题第一问灵敏度分析")
    parser.add_argument("--input", type=Path, default=q1.DEFAULT_INPUT)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=q1.OUTPUT_DIR / "sensitivity",
    )
    parser.add_argument("--mip-gap", type=float, default=1.0e-8)
    parser.add_argument("--time-limit", type=float, default=None)
    parser.add_argument("--solver-python", type=Path, default=None)
    args = parser.parse_args()

    base_df = q1.load_day_data(args.input)

    rows: list[dict[str, object]] = []

    def add(
        group: str,
        name: str,
        *,
        df: pd.DataFrame | None = None,
        eta: float = q1.DEFAULT_ETA,
        soc_min: float = q1.DEFAULT_SOC_MIN,
        soc_max: float = q1.DEFAULT_SOC_MAX,
        soc_initial: float = q1.DEFAULT_SOC_INITIAL,
        power_limit_kw: float = q1.DEFAULT_POWER_LIMIT_KW,
    ) -> None:
        rows.append(
            run_case(
                base_df if df is None else df,
                scenario_group=group,
                scenario_name=name,
                eta=eta,
                soc_min=soc_min,
                soc_max=soc_max,
                soc_initial=soc_initial,
                power_limit_kw=power_limit_kw,
                mip_gap=args.mip_gap,
                time_limit=args.time_limit,
                solver_python=args.solver_python,
            )
        )

    add("基准", "基准情景")

    for value in (1200, 2400, 3600, 6000, 7200, 8400, 9600, 10800):
        add("初始和终端SOC", f"SOC={value}", soc_initial=float(value))

    for value in (0.70, 0.75, 0.80, 0.85, 0.90, 0.95, 1.00):
        add("充放电效率", f"eta={value:.2f}", eta=float(value))

    for value in (-0.20, -0.10, -0.05, 0.00, 0.05, 0.10, 0.20):
        add(
            "PV预测",
            f"PV={100 * (1 + value):.0f}%",
            df=q1.load_day_data(args.input, pv_factor=1.0 + value),
        )

    for value in (-0.20, -0.10, -0.05, 0.00, 0.05, 0.10, 0.20):
        add(
            "负载预测",
            f"load={100 * (1 + value):.0f}%",
            df=q1.load_day_data(args.input, load_factor=1.0 + value),
        )

    for value in (3000, 4000, 5000, 6000, 7000, 8000):
        add(
            "最大充放电功率",
            f"Pmax={value} kW",
            power_limit_kw=float(value),
        )

    for lower, upper in (
        (2400, 9600),
        (1800, 10200),
        (1200, 10800),
        (600, 11400),
        (0, 12000),
    ):
        add(
            "储能可用SOC区间",
            f"{lower}-{upper} kWh",
            soc_min=float(lower),
            soc_max=float(upper),
        )

    combinations = (
        ("组合情景", "load+5%,PV-5%", 1.05, 0.95, 0.90),
        ("组合情景", "load+10%,PV-10%", 1.10, 0.90, 0.90),
        ("组合情景", "load+5%,PV-5%,eta=0.85", 1.05, 0.95, 0.85),
        ("组合情景", "load-5%,PV+5%", 0.95, 1.05, 0.90),
    )
    for _, name, load_factor, pv_factor, eta in combinations:
        add(
            "组合情景",
            name,
            df=q1.load_day_data(
                args.input,
                load_factor=load_factor,
                pv_factor=pv_factor,
            ),
            eta=eta,
        )

    frame = pd.DataFrame(rows)
    base_cost = float(
        frame.loc[frame["scenario_name"] == "基准情景", "cost"].iloc[0]
    )
    frame["delta_cost"] = frame["cost"] - base_cost
    frame["cost_change_pct"] = frame["delta_cost"] / base_cost * 100.0

    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.output_dir / "sensitivity_results.csv"
    xlsx_path = args.output_dir / "sensitivity_results.xlsx"
    md_path = args.output_dir / "sensitivity_summary.md"
    frame.to_csv(csv_path, index=False, encoding="utf-8-sig")
    write_results_workbook(frame, xlsx_path)
    md_path.write_text(
        build_summary_markdown(frame, base_cost),
        encoding="utf-8",
    )

    print(f"灵敏度结果: {csv_path}")
    print(f"灵敏度工作簿: {xlsx_path}")
    print(f"灵敏度报告: {md_path}")
    print(f"基准购电费: {base_cost:.6f} 元")
    print(frame[["scenario_group", "scenario_name", "cost_change_pct"]].to_string(index=False))


if __name__ == "__main__":
    main()
