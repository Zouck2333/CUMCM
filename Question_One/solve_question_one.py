from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "C题" / "附件" / "附件1.xlsx"
DEFAULT_TEMPLATE = ROOT / "C题" / "附件" / "附件5" / "result1.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parent / "output"

DELTA_H = 1.0 / 6.0
DEFAULT_ETA = 0.9
DEFAULT_SOC_MIN = 1200.0
DEFAULT_SOC_MAX = 10800.0
DEFAULT_SOC_INITIAL = 6000.0
DEFAULT_POWER_LIMIT_KW = 5000.0


def load_day_data(
    path: Path,
    *,
    load_factor: float = 1.0,
    pv_factor: float = 1.0,
) -> pd.DataFrame:
    df = pd.read_excel(path)
    required = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    missing = [name for name in required if name not in df.columns]
    if missing:
        raise ValueError(f"附件1缺少列: {missing}")
    if len(df) != 144:
        raise ValueError(f"附件1应有144个时段，实际为{len(df)}")

    out = df[required].copy()
    out["电价"] = pd.to_numeric(out["电价"], errors="raise")
    out["小区负载"] = pd.to_numeric(out["小区负载"], errors="raise")
    out["光伏发电预测功率"] = pd.to_numeric(
        out["光伏发电预测功率"], errors="raise"
    )
    out["负载电量"] = out["小区负载"] * DELTA_H * load_factor
    out["光伏电量"] = (
        out["光伏发电预测功率"] * DELTA_H * pv_factor
    ).clip(lower=0.0)
    return out


def _solve_stage(
    df: pd.DataFrame,
    *,
    method: str,
    objective: str,
    eta: float,
    soc_min: float,
    soc_max: float,
    soc_initial: float,
    soc_terminal: float,
    power_limit_kw: float,
    mip_gap: float,
    time_limit: float | None,
    solver_python: Path | None,
    cost_upper: float | None = None,
    throughput_upper: float | None = None,
) -> tuple[pd.DataFrame, dict[str, float | int | str]]:
    n_time = len(df)
    price = df["电价"].to_numpy(dtype=float)
    load = df["负载电量"].to_numpy(dtype=float)
    pv = df["光伏电量"].to_numpy(dtype=float)

    payload = {
        "price": price.tolist(),
        "load": load.tolist(),
        "pv": pv.tolist(),
        "delta_h": DELTA_H,
        "eta": eta,
        "soc_min": soc_min,
        "soc_max": soc_max,
        "soc_initial": soc_initial,
        "soc_terminal": soc_terminal,
        "power_limit_kw": power_limit_kw,
        "method": method,
        "objective": objective,
        "mip_gap": mip_gap,
        "time_limit": time_limit,
        "cost_upper": cost_upper,
        "throughput_upper": throughput_upper,
    }
    worker = Path(__file__).resolve().parent / "solver_worker.py"
    solver_executable = find_solver_python(solver_python)
    with tempfile.TemporaryDirectory(prefix="cumcm_optimized_") as temp_dir:
        input_json = Path(temp_dir) / "input.json"
        output_json = Path(temp_dir) / "output.json"
        input_json.write_text(json.dumps(payload), encoding="utf-8")
        completed = subprocess.run(
            [str(solver_executable), str(worker), str(input_json), str(output_json)],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        if completed.returncode != 0:
            message = completed.stderr.strip() or completed.stdout.strip()
            raise RuntimeError(f"{method.upper()}求解进程失败: {message}")
        solved = json.loads(output_json.read_text(encoding="utf-8"))

    grid = np.maximum(0.0, np.asarray(solved["grid"], dtype=float))
    charge = np.maximum(0.0, np.asarray(solved["charge"], dtype=float))
    discharge = np.maximum(
        0.0, np.asarray(solved["discharge"], dtype=float)
    )
    curtailment = np.maximum(
        0.0, np.asarray(solved["curtailment"], dtype=float)
    )
    soc = np.asarray(solved["soc"], dtype=float)

    records = []
    for t in range(n_time):
        e_before = soc_initial if t == 0 else soc[t - 1]
        records.append(
            {
                "时段": t + 1,
                "时间": df.iloc[t]["时间"],
                "电价": price[t],
                "负载电量_kWh": load[t],
                "光伏电量_kWh": pv[t],
                "计划购电量_kWh": grid[t],
                "充电量_kWh": charge[t],
                "放电量_kWh": discharge[t],
                "弃光或富余电量_kWh": curtailment[t],
                "时段初储电量_kWh": e_before,
                "时段末储电量_kWh": soc[t],
                "购电费用": price[t] * grid[t],
            }
        )

    info: dict[str, float | int | str] = {
        "model_type": method.upper(),
        "objective": objective,
        "solver": str(solved["solver"]),
        "solver_python": str(solver_executable),
        "solver_status": int(solved["status"]),
        "solver_message": str(solved["message"]),
        "objective_value": float(solved["objective_value"]),
        "total_cost": float(price @ grid),
        "throughput": float(charge.sum() + discharge.sum()),
        "peak_kw": float(np.max(grid / DELTA_H)),
    }
    if method == "milp":
        info["mip_gap"] = float(solved["mip_gap"])
        info["mip_node_count"] = int(solved["mip_node_count"])
    else:
        info["iterations"] = int(solved["iterations"])
    return pd.DataFrame(records), info


def find_solver_python(requested: Path | None) -> Path:
    candidates: list[Path] = []
    if requested is not None:
        candidates.append(requested)
    candidates.append(Path(sys.executable))
    path_python = shutil.which("python")
    if path_python:
        candidates.append(Path(path_python))

    checked: set[str] = set()
    for candidate in candidates:
        if not candidate.exists():
            continue
        resolved = candidate.resolve()
        key = str(resolved).lower()
        if key in checked:
            continue
        checked.add(key)
        probe = subprocess.run(
            [str(resolved), "-c", "from scipy.optimize import linprog, milp"],
            check=False,
            capture_output=True,
            text=True,
        )
        if probe.returncode == 0:
            return resolved
    raise RuntimeError(
        "未找到同时支持 scipy.optimize.linprog 和 milp 的 Python；"
        "请使用 --solver-python 指定求解器环境。"
    )


def solve_dispatch(
    df: pd.DataFrame,
    *,
    method: str,
    objective_mode: str,
    eta: float,
    soc_min: float,
    soc_max: float,
    soc_initial: float,
    soc_terminal: float,
    power_limit_kw: float,
    mip_gap: float,
    time_limit: float | None,
    solver_python: Path | None = None,
) -> tuple[pd.DataFrame, dict[str, object]]:
    if method == "lp":
        result, info = _solve_stage(
            df,
            method=method,
            objective="cost",
            eta=eta,
            soc_min=soc_min,
            soc_max=soc_max,
            soc_initial=soc_initial,
            soc_terminal=soc_terminal,
            power_limit_kw=power_limit_kw,
            mip_gap=mip_gap,
            time_limit=time_limit,
            solver_python=solver_python,
        )
        info["lexicographic"] = False
        return result, info

    if objective_mode == "cost":
        result, info = _solve_stage(
            df,
            method=method,
            objective="cost",
            eta=eta,
            soc_min=soc_min,
            soc_max=soc_max,
            soc_initial=soc_initial,
            soc_terminal=soc_terminal,
            power_limit_kw=power_limit_kw,
            mip_gap=mip_gap,
            time_limit=time_limit,
            solver_python=solver_python,
        )
        info["lexicographic"] = False
        return result, info

    cost_result, cost_info = _solve_stage(
        df,
        method=method,
        objective="cost",
        eta=eta,
        soc_min=soc_min,
        soc_max=soc_max,
        soc_initial=soc_initial,
        soc_terminal=soc_terminal,
        power_limit_kw=power_limit_kw,
        mip_gap=mip_gap,
        time_limit=time_limit,
        solver_python=solver_python,
    )
    cost_opt = float(cost_info["total_cost"])
    cost_tol = max(1e-5, abs(cost_opt) * 1e-9)

    throughput_result, throughput_info = _solve_stage(
        df,
        method=method,
        objective="throughput",
        eta=eta,
        soc_min=soc_min,
        soc_max=soc_max,
        soc_initial=soc_initial,
        soc_terminal=soc_terminal,
        power_limit_kw=power_limit_kw,
        mip_gap=mip_gap,
        time_limit=time_limit,
        solver_python=solver_python,
        cost_upper=cost_opt + cost_tol,
    )
    throughput_opt = float(throughput_info["throughput"])
    throughput_tol = max(1e-5, abs(throughput_opt) * 1e-9)

    peak_result, peak_info = _solve_stage(
        df,
        method=method,
        objective="peak",
        eta=eta,
        soc_min=soc_min,
        soc_max=soc_max,
        soc_initial=soc_initial,
        soc_terminal=soc_terminal,
        power_limit_kw=power_limit_kw,
        mip_gap=mip_gap,
        time_limit=time_limit,
        solver_python=solver_python,
        cost_upper=cost_opt + cost_tol,
        throughput_upper=throughput_opt + throughput_tol,
    )
    peak_info["lexicographic"] = True
    peak_info["lexicographic_stages"] = {
        "cost": cost_opt,
        "throughput": throughput_opt,
        "peak_kw": float(peak_info["peak_kw"]),
        "cost_tolerance": cost_tol,
        "throughput_tolerance": throughput_tol,
    }
    peak_info["initial_cost_solution"] = {
        "cost": cost_opt,
        "throughput": float(cost_info["throughput"]),
        "peak_kw": float(cost_info["peak_kw"]),
    }
    peak_info["throughput_solution"] = {
        "cost": float(throughput_info["total_cost"]),
        "throughput": throughput_opt,
        "peak_kw": float(throughput_info["peak_kw"]),
    }
    return peak_result, peak_info


def validate_solution(
    result: pd.DataFrame,
    energy_limit: float,
    eta: float,
    soc_min: float,
    soc_max: float,
    soc_terminal: float,
) -> dict[str, float]:
    balance = (
        result["计划购电量_kWh"]
        + result["光伏电量_kWh"]
        + result["放电量_kWh"]
        - result["负载电量_kWh"]
        - result["充电量_kWh"]
        - result["弃光或富余电量_kWh"]
    ).abs().max()
    soc_error = (
        result["时段初储电量_kWh"]
        + eta * result["充电量_kWh"]
        - result["放电量_kWh"] / eta
        - result["时段末储电量_kWh"]
    ).abs().max()
    simultaneous = int(
        (
            (result["充电量_kWh"] > 1e-7)
            & (result["放电量_kWh"] > 1e-7)
        ).sum()
    )
    power_violations = int(
        (
            (result["充电量_kWh"] > energy_limit + 1e-7)
            | (result["放电量_kWh"] > energy_limit + 1e-7)
        ).sum()
    )
    curtailment_violations = int(
        (
            (result["弃光或富余电量_kWh"] < -1e-7)
            | (result["弃光或富余电量_kWh"] > result["光伏电量_kWh"] + 1e-7)
        ).sum()
    )
    soc_lower_violations = int((result["时段末储电量_kWh"] < soc_min - 1e-7).sum())
    soc_upper_violations = int((result["时段末储电量_kWh"] > soc_max + 1e-7).sum())
    terminal_soc_error = abs(float(result.iloc[-1]["时段末储电量_kWh"]) - soc_terminal)
    simultaneous_kwh = float(
        np.minimum(result["充电量_kWh"], result["放电量_kWh"]).sum()
    )
    return {
        "total_cost": float(result["购电费用"].sum()),
        "total_grid_kwh": float(result["计划购电量_kWh"].sum()),
        "total_charge_kwh": float(result["充电量_kWh"].sum()),
        "total_discharge_kwh": float(result["放电量_kWh"].sum()),
        "max_balance_error": float(balance),
        "max_soc_error": float(soc_error),
        "initial_soc": float(result.iloc[0]["时段初储电量_kWh"]),
        "terminal_soc": float(result.iloc[-1]["时段末储电量_kWh"]),
        "min_soc": float(result["时段末储电量_kWh"].min()),
        "max_soc": float(result["时段末储电量_kWh"].max()),
        "max_charge_kwh": float(result["充电量_kWh"].max()),
        "max_discharge_kwh": float(result["放电量_kWh"].max()),
        "power_violations": power_violations,
        "curtailment_violations": curtailment_violations,
        "soc_lower_bound_violations": soc_lower_violations,
        "soc_upper_bound_violations": soc_upper_violations,
        "terminal_soc_error": terminal_soc_error,
        "simultaneous_charge_discharge_periods": simultaneous,
        "simultaneous_charge_discharge_kwh": simultaneous_kwh,
    }


def validate_parameters(
    *,
    eta: float,
    soc_min: float,
    soc_max: float,
    soc_initial: float,
    soc_terminal: float,
    power_limit_kw: float,
    load_factor: float,
    pv_factor: float,
    mip_gap: float,
    time_limit: float | None,
) -> None:
    if not 0.0 < eta <= 1.0:
        raise ValueError("eta必须满足0 < eta <= 1")
    if soc_min >= soc_max:
        raise ValueError("soc-min必须小于soc-max")
    if not soc_min <= soc_initial <= soc_max:
        raise ValueError("soc-initial必须位于SOC上下限内")
    if not soc_min <= soc_terminal <= soc_max:
        raise ValueError("soc-terminal必须位于SOC上下限内")
    if power_limit_kw <= 0.0:
        raise ValueError("power-limit-kw必须大于0")
    if load_factor < 0.0 or pv_factor < 0.0:
        raise ValueError("负载和光伏缩放系数不能为负数")
    if mip_gap < 0.0:
        raise ValueError("mip-gap不能为负数")
    if time_limit is not None and time_limit <= 0.0:
        raise ValueError("time-limit必须大于0")


def fill_template(
    template_path: Path,
    output_path: Path,
    result: pd.DataFrame,
    *,
    soc_initial: float,
    soc_terminal: float,
) -> None:
    workbook = openpyxl.load_workbook(template_path)
    plan_sheet = workbook["计划购电量"]
    for row_idx, value in enumerate(result["计划购电量_kWh"], start=2):
        plan_sheet.cell(row=row_idx, column=2).value = round(
            float(value), 4
        )

    battery_sheet = workbook["充放电量"]
    for block in range(6):
        start = block * 24
        stop = start + 24
        battery_sheet.cell(row=block + 2, column=2).value = round(
            float(result.iloc[start:stop]["充电量_kWh"].sum()), 4
        )
        battery_sheet.cell(row=block + 2, column=3).value = round(
            float(result.iloc[start:stop]["放电量_kWh"].sum()), 4
        )
    battery_sheet.cell(row=2, column=5).value = round(soc_initial, 4)
    battery_sheet.cell(row=3, column=5).value = round(soc_terminal, 4)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="C题第一问参数化LP/MILP求解程序"
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--method", choices=("lp", "milp"), default="milp")
    parser.add_argument(
        "--objective-mode",
        choices=("auto", "cost", "lexicographic"),
        default="auto",
    )
    parser.add_argument("--eta", type=float, default=DEFAULT_ETA)
    parser.add_argument("--soc-min", type=float, default=DEFAULT_SOC_MIN)
    parser.add_argument("--soc-max", type=float, default=DEFAULT_SOC_MAX)
    parser.add_argument(
        "--soc-initial", type=float, default=DEFAULT_SOC_INITIAL
    )
    parser.add_argument("--soc-terminal", type=float, default=None)
    parser.add_argument(
        "--power-limit-kw", type=float, default=DEFAULT_POWER_LIMIT_KW
    )
    parser.add_argument("--load-factor", type=float, default=1.0)
    parser.add_argument("--pv-factor", type=float, default=1.0)
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--detail", type=Path, default=None)
    parser.add_argument("--mip-gap", type=float, default=1.0e-8)
    parser.add_argument("--time-limit", type=float, default=None)
    parser.add_argument("--solver-python", type=Path, default=None)
    args = parser.parse_args()

    objective_mode = args.objective_mode
    if objective_mode == "auto":
        objective_mode = "cost"
    if args.method == "lp" and objective_mode == "lexicographic":
        raise ValueError("LP does not support lexicographic optimization.")

    soc_terminal = (
        args.soc_initial
        if args.soc_terminal is None
        else args.soc_terminal
    )
    validate_parameters(
        eta=args.eta,
        soc_min=args.soc_min,
        soc_max=args.soc_max,
        soc_initial=args.soc_initial,
        soc_terminal=soc_terminal,
        power_limit_kw=args.power_limit_kw,
        load_factor=args.load_factor,
        pv_factor=args.pv_factor,
        mip_gap=args.mip_gap,
        time_limit=args.time_limit,
    )
    output_path = (
        args.output
        if args.output is not None
        else OUTPUT_DIR / f"result1_{args.method}.xlsx"
    )
    detail_path = (
        args.detail
        if args.detail is not None
        else OUTPUT_DIR / f"question_one_detail_{args.method}.csv"
    )
    df = load_day_data(
        args.input,
        load_factor=args.load_factor,
        pv_factor=args.pv_factor,
    )
    result, solver_info = solve_dispatch(
        df,
        method=args.method,
        objective_mode=objective_mode,
        eta=args.eta,
        soc_min=args.soc_min,
        soc_max=args.soc_max,
        soc_initial=args.soc_initial,
        soc_terminal=soc_terminal,
        power_limit_kw=args.power_limit_kw,
        mip_gap=args.mip_gap,
        time_limit=args.time_limit,
        solver_python=args.solver_python,
    )
    checks = validate_solution(
        result,
        energy_limit=args.power_limit_kw * DELTA_H,
        eta=args.eta,
        soc_min=args.soc_min,
        soc_max=args.soc_max,
        soc_terminal=soc_terminal,
    )
    detail_path.parent.mkdir(parents=True, exist_ok=True)
    result.to_csv(detail_path, index=False, encoding="utf-8-sig")
    fill_template(
        args.template,
        output_path,
        result,
        soc_initial=args.soc_initial,
        soc_terminal=soc_terminal,
    )

    print(f"C题第一问{args.method.upper()}求解完成")
    print(f"输出模板: {output_path}")
    print(f"详细结果: {detail_path}")
    for key, value in solver_info.items():
        print(f"{key}: {json.dumps(value, ensure_ascii=False)}")
    for key, value in checks.items():
        print(f"{key}: {value:.6f}")


if __name__ == "__main__":
    main()
