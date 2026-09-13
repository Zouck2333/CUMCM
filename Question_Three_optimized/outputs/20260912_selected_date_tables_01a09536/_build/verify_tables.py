import hashlib
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

HERE=Path(__file__).resolve().parent
data=json.loads((HERE/'tables_data.json').read_text(encoding='utf-8'))
expected=json.loads((HERE/'expected_cells.json').read_text(encoding='utf-8'))
book=HERE.parent/'第三问优化结果_表1表2表3.xlsx'
values=load_workbook(book,data_only=True)
formulas=load_workbook(book,data_only=False)
assert len(values.worksheets)==4
for e in expected:
    v=values[e['sheet']][e['cell']].value
    assert isinstance(v,(int,float)) and abs(v-e['value'])<1e-6,(e,v)
    assert formulas[e['sheet']][e['cell']].data_type=='f',e
for row in values['所选日期明细'].iter_rows(min_row=5,values_only=True):
    assert len(row)==18 and all(v is not None for v in row)
for s in values:
    for row in s:
        for c in row:
            assert c.data_type!='e',(s.title,c.coordinate,c.value)
for i,d in enumerate(data['days']):
    for name in ['表1 购电量与费用','表2 储能充放电']:
        assert values[name].cell(3+i*11,1).value.strftime('%Y-%m-%d')==d['date']
    assert values['表3 紧急购电'].cell(5,1+i*2).value.strftime('%Y-%m-%d')==d['date']
docx=HERE.parent/'第三问优化结果_表1表2表3.docx'
doc=Document(docx)
assert len(doc.tables)==9
for i,d in enumerate(data['days']):
    t1,t2=doc.tables[i*2:i*2+2]
    for j,p in enumerate(d['selected']):
        assert t1.cell(1+j//3,(j%3)*2).text==p['period']
        assert t1.cell(1+j//3,(j%3)*2+1).text==f"{p['energy']:.4f}"
    assert t1.cell(3,2).text==f"{d['totals']['final_purchase_kwh']:.4f}"
    assert t1.cell(3,5).text==f"{d['totals']['total_cost_yuan']:.2f}"
    for j,b in enumerate(d['blocks']):
        r,c=1+j//2,(j%2)*3
        assert t2.cell(r,c).text==b['period']
        assert t2.cell(r,c+1).text==f"{b['charge']:.4f}"
        assert t2.cell(r,c+2).text==f"{b['discharge']:.4f}"
    assert t2.cell(4,2).text==f"{d['soc_start']:.4f}"
    assert t2.cell(4,5).text==f"{d['soc_end']:.4f}"
    for j,g in enumerate(d['emergency']):
        assert doc.tables[8].cell(j+2,i*2).text==g['period']
        assert doc.tables[8].cell(j+2,i*2+1).text==f"{g['energy']:.4f}"
pdf=PdfReader(HERE/'word_render.pdf')
assert len(pdf.pages)==5,len(pdf.pages)
for i,p in enumerate(pdf.pages):
    text=p.extract_text()
    assert len(text)>300,(i,len(text))
report={'status':'passed','source_detail_rows':576,'excel_result_cells_checked':len(expected),'excel_formula_error_count':0,'word_tables':9,'word_pages':len(pdf.pages),'emergency_intervals':[len(d['emergency']) for d in data['days']],'render_engine':'Microsoft Word + render_docx.py / bundled Poppler','source_hashes':data['source_hashes'],'output_hashes':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [book,docx]}}
(HERE/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
