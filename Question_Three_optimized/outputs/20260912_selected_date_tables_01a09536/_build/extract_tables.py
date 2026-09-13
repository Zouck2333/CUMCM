"""Read the verified optimization outputs and reconcile requested paper tables."""
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SOURCE = ROOT / 'output_optimized'
DATES = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']

def close(a, b):
    assert abs(a - b) < 1e-6, (a, b)

def clock(t):
    return f'{t//6}:{t%6*10:02d}'

def main():
    with (SOURCE / 'question_three_detail.csv').open(encoding='utf-8-sig', newline='') as f:
        records = [r for r in csv.DictReader(f) if r['date'] in DATES]
    for r in records:
        for k in r:
            if k not in ('date', 'time_period'):
                r[k] = int(r[k]) if k == 't' else float(r[k])
    checkpoint = {}
    with (SOURCE / 'checkpoint_rolling.jsonl').open(encoding='utf-8') as f:
        for line in f:
            r = json.loads(line)
            if r['date'] in DATES:
                checkpoint[r['date']] = r
    wb = load_workbook(SOURCE / 'result3.xlsx', data_only=True, read_only=True)
    adjusted = {r[0].strftime('%Y-%m-%d'): r for r in wb.worksheets[1].iter_rows(min_row=2, values_only=True) if r[0]}
    storage = {}
    day = None
    for row in wb.worksheets[2].iter_rows(min_row=2, values_only=True):
        if row[0]:
            day = row[0].strftime('%Y-%m-%d')
        storage.setdefault(day, []).append(row)
    emergency = {}
    for row in wb.worksheets[3].iter_rows(min_row=2, values_only=True):
        if row[0]:
            day = row[0].strftime('%Y-%m-%d')
        if row[1]:
            emergency.setdefault(day, []).append([row[1], row[2]])
    days = []
    for di, date in enumerate(DATES):
        rows = [r for r in records if r['date'] == date]
        assert [r['t'] for r in rows] == list(range(144))
        cp = checkpoint[date]
        for r in rows:
            for field, key in [('final_purchase_kwh', 'final_purchase'), ('initial_purchase_kwh', 'initial_purchase'), ('charge_kwh', 'charge'), ('discharge_kwh', 'discharge'), ('emergency_kwh', 'emergency')]:
                close(r[field], cp[key][r['t']])
            close(r['total_cost_yuan'], r['plan_cost_yuan'] + r['adjustment_cost_yuan'] + r['emergency_cost_yuan'])
            close(r['soc_after_kwh'], r['soc_before_kwh'] + .9*r['charge_kwh'] - r['discharge_kwh']/.9)
        total = {k: sum(r[k] for r in rows) for k in ['final_purchase_kwh', 'emergency_kwh', 'plan_cost_yuan', 'adjustment_cost_yuan', 'emergency_cost_yuan', 'total_cost_yuan']}
        close(total['total_cost_yuan'], cp['total_cost'])
        close(total['final_purchase_kwh'], adjusted[date][145])
        close(total['total_cost_yuan'], adjusted[date][146])
        selected = []
        for hour in [10, 12, 14, 16, 18, 20]:
            t = hour * 6
            close(rows[t]['final_purchase_kwh'], adjusted[date][t+1])
            selected.append({'period': f'{clock(t)}-{clock(t+1)}', 'energy': rows[t]['final_purchase_kwh'], 'detail_row': 5+di*144+t})
        blocks = []
        for b in range(6):
            start, end = b*24, (b+1)*24
            c = sum(r['charge_kwh'] for r in rows[start:end])
            d = sum(r['discharge_kwh'] for r in rows[start:end])
            close(c, storage[date][b][2])
            close(d, storage[date][b][3])
            blocks.append({'period': f'{clock(start)}-{clock(end)}', 'charge': c, 'discharge': d, 'first_row': 5+di*144+start, 'last_row': 5+di*144+end-1})
        close(rows[0]['soc_before_kwh'], storage[date][0][5])
        close(rows[-1]['soc_after_kwh'], storage[date][1][5])
        groups = []
        start = None
        for t in range(145):
            if t < 144 and rows[t]['emergency_kwh'] > 1e-8:
                if start is None:
                    start = t
            elif start is not None:
                groups.append({'period': f'{clock(start)}-{clock(t)}', 'energy': sum(r['emergency_kwh'] for r in rows[start:t]), 'first_row': 5+di*144+start, 'last_row': 5+di*144+t-1})
                start = None
        assert len(groups) == len(emergency[date])
        for g, (label, amount) in zip(groups, emergency[date]):
            assert g['period'] == label, (g['period'], label)
            close(g['energy'], amount)
        close(sum(g['energy'] for g in groups), total['emergency_kwh'])
        days.append({'date': date, 'selected': selected, 'blocks': blocks, 'emergency': groups, 'totals': total, 'soc_start': rows[0]['soc_before_kwh'], 'soc_end': rows[-1]['soc_after_kwh'], 'first_row': 5+di*144, 'last_row': 148+di*144})
    wb.close()
    result = {'days': days, 'records': records, 'source_hashes': {name: hashlib.sha256((SOURCE/name).read_bytes()).hexdigest() for name in ['question_three_detail.csv', 'checkpoint_rolling.jsonl', 'result3.xlsx']}, 'validation': {'status': 'passed', 'detail_rows': len(records), 'source_workbook': 'result3.xlsx', 'cross_checked': ['checkpoint', 'final_purchase', 'daily_cost', 'battery_blocks', 'soc', 'emergency_intervals']}}
    (HERE / 'tables_data.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'status': 'passed', 'days': [{'date': d['date'], **d['totals'], 'emergency_groups': len(d['emergency'])} for d in days]}, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
