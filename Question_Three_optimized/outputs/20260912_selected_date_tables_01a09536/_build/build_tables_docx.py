"""Create editable Word tables following the supplied image layouts."""
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
DATA = json.loads((HERE/'tables_data.json').read_text(encoding='utf-8'))
DOC = Document()
section = DOC.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11.69)
section.page_height = Inches(8.27)
section.top_margin = section.bottom_margin = Inches(.5)
section.left_margin = section.right_margin = Inches(.6)
for name in ['Normal', 'Title', 'Heading 1', 'Caption']:
    style = DOC.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.space_after = Pt(0)

def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, style=None):
    p = DOC.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.name = 'Times New Roman'
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    return p

def table(rows, cols, height=31, size=12):
    t = DOC.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col in t.columns:
        col.width = Inches(10.49/cols)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:'+edge)
        for k,v in {'val':'single','sz':'8','color':'000000'}.items():
            e.set(qn('w:'+k),v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        e = OxmlElement('w:'+edge)
        e.set(qn('w:w'),'35')
        e.set(qn('w:type'),'dxa')
        margins.append(e)
    t._tbl.tblPr.append(margins)
    for row in t.rows:
        row.height = Pt(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraphs[0].paragraph_format.space_before = Pt(0)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.line_spacing = Pt(size+2)
            snap = OxmlElement('w:snapToGrid')
            snap.set(qn('w:val'), '0')
            c.paragraphs[0]._p.get_or_add_pPr().append(snap)
    return t

def put(t, row, col, val, size=12):
    c=t.cell(row,col)
    c.text=''
    p=c.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=Pt(size+2)
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)
    r=p.add_run(val if isinstance(val,str) else f'{val:.4f}')
    r.font.name='Times New Roman'
    r.font.size=Pt(size)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')

for i,d in enumerate(DATA['days']):
    if i:
        DOC.add_page_break()
    year,month,day=[int(v) for v in d['date'].split('-')]
    para(f'{year}年{month}月{day}日',18,True,after=12,style='Title')
    para('表1  微网在指定时间段的购电量及全天的购电量和购电费',14,True,after=9)
    t=table(4,6,height=29,size=12)
    for c,v in enumerate(['时间段','购电量','时间段','购电量','时间段','购电量']): put(t,0,c,v)
    for j,p in enumerate(d['selected']):
        rr,cc=1+j//3,(j%3)*2
        put(t,rr,cc,p['period']); put(t,rr,cc+1,p['energy'])
    t.cell(3,0).merge(t.cell(3,1)); t.cell(3,3).merge(t.cell(3,4))
    put(t,3,0,'全天购电量'); put(t,3,2,d['totals']['final_purchase_kwh'])
    put(t,3,3,'全天购电费'); put(t,3,5,f"{d['totals']['total_cost_yuan']:.2f}")
    para('单位：购电量为kWh，购电费为元。购电量为最终正常购电量，紧急购电量见表3。',10,align=WD_ALIGN_PARAGRAPH.LEFT,before=6,after=3)
    para('全天购电费＝计划购电费＋调整费＋紧急购电费。',10,align=WD_ALIGN_PARAGRAPH.LEFT,after=11)
    para('表2  储能设备在指定时间段的充放电量及0:00和24:00的储电量',14,True,after=9)
    t=table(5,6,height=28,size=12)
    for c,v in enumerate(['时间段','充电量','放电量','时间段','充电量','放电量']): put(t,0,c,v)
    for j,b in enumerate(d['blocks']):
        rr,cc=1+j//2,(j%2)*3
        for k,v in enumerate([b['period'],b['charge'],b['discharge']]):put(t,rr,cc+k,v)
    t.cell(4,0).merge(t.cell(4,1));t.cell(4,3).merge(t.cell(4,4))
    put(t,4,0,'0:00 储电量');put(t,4,2,d['soc_start'])
    put(t,4,3,'24:00 储电量');put(t,4,5,d['soc_end'])
    para('单位：kWh。充放电量按微网侧计量，为各4小时时段实际执行量之和；储电量按电池内部SOC计量。',10,align=WD_ALIGN_PARAGRAPH.LEFT,before=6,after=0)

DOC.add_page_break()
para('表3  微网在指定日期的紧急购电量',16,True,before=0,after=9,style='Title')
para('单位：kWh',10,align=WD_ALIGN_PARAGRAPH.RIGHT,after=7)
n=max(len(d['emergency']) for d in DATA['days'])
t=table(n+2,8,height=19,size=11)
for i,d in enumerate(DATA['days']):
    t.cell(0,2*i).merge(t.cell(0,2*i+1))
    put(t,0,2*i,'.'.join(str(int(v)) for v in d['date'].split('-')),12)
    put(t,1,2*i,'时间段',11);put(t,1,2*i+1,'购电量',11)
    for j,g in enumerate(d['emergency']):
        put(t,j+2,2*i,g['period'],11);put(t,j+2,2*i+1,g['energy'],11)
for row in t.rows[:2]:
    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
para('注：相邻且均发生紧急购电的10分钟时段合并列示；空白表示该日期已无更多紧急购电时段。',10,align=WD_ALIGN_PARAGRAPH.LEFT,before=7,after=0)
DOC.core_properties.title='第三问优化结果指定日期表格'
DOC.core_properties.subject='2025年3月20日、6月21日、9月23日、12月21日的表1、表2和表3'
DOC.core_properties.author=''
# Remove inherited title paragraph rules from the bundled default template.
for root in [DOC.styles.element, DOC.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
target=HERE.parent/'第三问优化结果_表1表2表3.docx'
DOC.save(target)
print(json.dumps({'path':str(target),'tables':len(DOC.tables)},ensure_ascii=False))
