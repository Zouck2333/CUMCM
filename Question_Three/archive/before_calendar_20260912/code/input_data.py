"""Read and validate the source workbooks for CUMCM problem C, question three.

Power observations in attachments 1 and 2 are recorded at the *right endpoint*
of each ten-minute interval.  The arrays returned by this module use zero-based
interval positions, so column B (00:10) is interval [00:00, 00:10).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
import math

import numpy as np

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    if exc.name != "openpyxl":
        raise
    # The desktop runtime ships openpyxl separately from the system Python
    # that supplies SciPy. Append its pure-Python packages only when needed;
    # never put them ahead of the active interpreter's NumPy/SciPy packages.
    import sys

    bundled_packages = (
        Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime"
        / "dependencies" / "python" / "Lib" / "site-packages"
    )
    if not (bundled_packages / "openpyxl").is_dir():
        raise ModuleNotFoundError(
            "需要 openpyxl；当前 Python 未安装，且找不到 Codex 的捆绑依赖"
        ) from exc
    sys.path.append(str(bundled_packages))
    from openpyxl import load_workbook


INTERVALS_PER_DAY = 144
HOURS_PER_FORECAST = 24
ISSUE_HOURS = (0, 6, 12, 18)
HOURS_PER_INTERVAL = 1.0 / 6.0
FIRST_DATE = date(2025, 1, 1)
LAST_DATE = date(2025, 12, 31)
DAYS = (LAST_DATE - FIRST_DATE).days + 1


@dataclass(frozen=True)
class InputData:
    """Validated inputs in chronological and physical interval order.

    ``prices[t]`` is yuan/kWh. ``load_kwh[d,t]`` and ``pv_kwh[d,t]``
    are energy during one ten-minute interval. ``pv_forecast_kw[d,k,j]``
    is the point power forecast for the (j+1)-th whole hour after issue k;
    this loader deliberately leaves the model's interpolation to the solver.
    """

    prices: np.ndarray
    load_kwh: np.ndarray
    pv_kwh: np.ndarray
    pv_forecast_kw: np.ndarray
    dates: list[date]
    template_path: Path


def _attachment_dir(base_dir: Path) -> Path:
    root = Path(base_dir).expanduser().resolve()
    for candidate in (root / "C题" / "附件", root / "附件", root):
        if all((candidate / f"附件{index}.xlsx").is_file() for index in (1, 2, 3)):
            return candidate
    raise FileNotFoundError(
        f"找不到附件1.xlsx、附件2.xlsx、附件3.xlsx；检查目录 {root}"
    )


def _number(value: object, context: str) -> float:
    if value is None or isinstance(value, bool):
        raise ValueError(f"{context} 缺少有效数值")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{context} 不是数值: {value!r}") from exc
    if not math.isfinite(result) or result < 0:
        raise ValueError(f"{context} 必须是非负有限数: {value!r}")
    return result


def _date(value: object, context: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            parts = value.strip().split("-")
            if len(parts) != 3 or not all(part.isdigit() for part in parts):
                raise ValueError("expected year-month-day")
            return date(*(int(part) for part in parts))
        except ValueError as exc:
            raise ValueError(f"{context} 日期无效: {value!r}") from exc
    raise ValueError(f"{context} 日期无效: {value!r}")


def _minute_of_day(value: object, context: str) -> int:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"{context} 时刻包含秒: {value!r}")
        return value.hour * 60 + value.minute
    if isinstance(value, str):
        label = value.strip()
        if label in {"0:00+1", "00:00+1", "24:00"}:
            return 1440
        pieces = label.split(":")
        if len(pieces) in (2, 3):
            try:
                hour, minute = int(pieces[0]), int(pieces[1])
                second = int(pieces[2]) if len(pieces) == 3 else 0
            except ValueError as exc:
                raise ValueError(f"{context} 时刻无效: {value!r}") from exc
            if 0 <= hour < 24 and 0 <= minute < 60 and second == 0:
                return hour * 60 + minute
    raise ValueError(f"{context} 时刻无效: {value!r}")


def _check_interval_headers(values: tuple[object, ...], context: str) -> None:
    if len(values) != INTERVALS_PER_DAY:
        raise ValueError(f"{context} 必须有144个时段标题")
    for t, value in enumerate(values):
        actual = _minute_of_day(value, f"{context} 第{t + 1}项")
        expected = (t + 1) * 10
        if actual != expected:
            raise ValueError(
                f"{context} 第{t + 1}项应为第{expected}分钟的右端点，实际为{value!r}"
            )


def _load_prices(path: Path) -> np.ndarray:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_row != INTERVALS_PER_DAY + 1 or sheet.max_column < 4:
            raise ValueError("附件1应有1行表头、144行价格和4列数据")
        rows = sheet.iter_rows(
            min_row=2, max_row=INTERVALS_PER_DAY + 1,
            min_col=1, max_col=4, values_only=True,
        )
        prices = np.empty(INTERVALS_PER_DAY, dtype=float)
        time_labels: list[object] = []
        for t, row in enumerate(rows):
            time_labels.append(row[0])
            prices[t] = _number(row[1], f"附件1 B{t + 2} 电价")
            _number(row[2], f"附件1 C{t + 2} 负载功率")
            _number(row[3], f"附件1 D{t + 2} 光伏功率")
        _check_interval_headers(tuple(time_labels), "附件1时段")
        return prices
    finally:
        workbook.close()


def _load_actuals(path: Path) -> tuple[list[date], np.ndarray, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        load_name = "小区负载"
        pv_name = "光伏发电实际功率"
        if load_name not in workbook or pv_name not in workbook:
            raise ValueError(f"附件2缺少工作表 {load_name!r} 或 {pv_name!r}")
        sheets = (workbook[load_name], workbook[pv_name])
        for sheet in sheets:
            if sheet.max_row != DAYS + 1 or sheet.max_column < INTERVALS_PER_DAY + 1:
                raise ValueError(f"附件2工作表{sheet.title!r}应有365天、每天144个功率点")
            headers = next(sheet.iter_rows(
                min_row=1, max_row=1, min_col=2,
                max_col=INTERVALS_PER_DAY + 1, values_only=True,
            ))
            _check_interval_headers(headers, f"附件2 {sheet.title} 时段")

        dates: list[date] = []
        actuals = [np.empty((DAYS, INTERVALS_PER_DAY), dtype=float) for _ in sheets]
        iterators = [sheet.iter_rows(
            min_row=2, max_row=DAYS + 1, min_col=1,
            max_col=INTERVALS_PER_DAY + 1, values_only=True,
        ) for sheet in sheets]
        for day_index, row_pair in enumerate(zip(*iterators, strict=True)):
            wanted_date = FIRST_DATE + timedelta(days=day_index)
            row_number = day_index + 2
            row_dates = [_date(row[0], f"附件2 {sheet.title} A{row_number}")
                         for sheet, row in zip(sheets, row_pair, strict=True)]
            if row_dates != [wanted_date, wanted_date]:
                raise ValueError(
                    f"附件2第{row_number}行日期应均为{wanted_date}，实际为{row_dates}"
                )
            dates.append(wanted_date)
            for array, sheet, row in zip(actuals, sheets, row_pair, strict=True):
                for t in range(INTERVALS_PER_DAY):
                    power_kw = _number(
                        row[t + 1], f"附件2 {sheet.title} 行{row_number} 时段{t}"
                    )
                    array[day_index, t] = power_kw * HOURS_PER_INTERVAL
        return dates, actuals[0], actuals[1]
    finally:
        workbook.close()


def _load_forecasts(path: Path, dates: list[date]) -> np.ndarray:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_row != 1 + DAYS * len(ISSUE_HOURS) or sheet.max_column < 26:
            raise ValueError("附件3应有365天×4次发布，每次24个整点预报")
        headers = next(sheet.iter_rows(
            min_row=1, max_row=1, min_col=1, max_col=26, values_only=True,
        ))
        expected_headers = tuple(f"预报{j}小时" for j in range(1, 25))
        if tuple(headers[2:]) != expected_headers:
            raise ValueError("附件3第3至26列应依次为预报1小时至预报24小时")

        forecasts = np.empty((DAYS, len(ISSUE_HOURS), HOURS_PER_FORECAST), dtype=float)
        rows = sheet.iter_rows(
            min_row=2, max_row=1 + DAYS * len(ISSUE_HOURS),
            min_col=1, max_col=26, values_only=True,
        )
        for index, row in enumerate(rows):
            day_index, issue_index = divmod(index, len(ISSUE_HOURS))
            row_number = index + 2
            date_cell = row[0]
            if issue_index == 0:
                found_date = _date(date_cell, f"附件3 A{row_number}")
                if found_date != dates[day_index]:
                    raise ValueError(
                        f"附件3 A{row_number} 日期应为{dates[day_index]}，实际为{found_date}"
                    )
            elif date_cell not in (None, ""):
                raise ValueError(f"附件3 A{row_number} 非首发布行应为空")
            issue_minute = _minute_of_day(row[1], f"附件3 B{row_number}")
            if issue_minute != ISSUE_HOURS[issue_index] * 60:
                raise ValueError(
                    f"附件3 B{row_number} 应为{ISSUE_HOURS[issue_index]}:00，实际为{row[1]!r}"
                )
            for hour_index in range(HOURS_PER_FORECAST):
                forecasts[day_index, issue_index, hour_index] = _number(
                    row[hour_index + 2],
                    f"附件3 行{row_number} 预报第{hour_index + 1}小时",
                )
        return forecasts
    finally:
        workbook.close()


def _check_template(path: Path, dates: list[date]) -> Path:
    if not path.is_file():
        raise FileNotFoundError(f"找不到第三问结果模板: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 4:
            raise ValueError("result3.xlsx 模板应有4个工作表")
        for sheet in workbook.worksheets[:2]:
            if sheet.max_row < 1 + (DAYS - 31) or sheet.max_column < 1 + INTERVALS_PER_DAY:
                raise ValueError(f"result3.xlsx 工作表{sheet.title!r}缺少2—12月数据区")
            first_date = _date(sheet.cell(2, 1).value, f"模板 {sheet.title} A2")
            last_date = _date(sheet.cell(DAYS - 31 + 1, 1).value,
                              f"模板 {sheet.title} A{DAYS - 31 + 1}")
            if first_date != dates[31] or last_date != dates[-1]:
                raise ValueError(f"result3.xlsx 工作表{sheet.title!r}日期范围与附件2不一致")
        return path
    finally:
        workbook.close()


def load_inputs(base_dir: Path) -> InputData:
    """Load original C-problem files from the workspace, ``C题``, or ``附件``.

    The array shapes are ``(144,)``, ``(365,144)``, ``(365,144)``, and
    ``(365,4,24)``. Dates are the 365 successive days of 2025. Fail fast on
    missing or malformed source data rather than replacing it with zeros.
    """

    attachment_dir = _attachment_dir(base_dir)
    prices = _load_prices(attachment_dir / "附件1.xlsx")
    dates, load_kwh, pv_kwh = _load_actuals(attachment_dir / "附件2.xlsx")
    pv_forecast_kw = _load_forecasts(attachment_dir / "附件3.xlsx", dates)
    template_path = _check_template(attachment_dir / "附件5" / "result3.xlsx", dates)
    return InputData(
        prices=prices,
        load_kwh=load_kwh,
        pv_kwh=pv_kwh,
        pv_forecast_kw=pv_forecast_kw,
        dates=dates,
        template_path=template_path,
    )
