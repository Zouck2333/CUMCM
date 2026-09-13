"""Build the three requested paper-table layouts from verified current outputs."""
from pathlib import Path
from datetime import datetime
import csv
import hashlib
import json
import math

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output"
DEST = OUT / "paper_tables/第二问_指定日期表1表2表3.docx"
DATES = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
HOURS = [10, 12, 14, 16, 18, 20]
checks = []


def check(label, condition):
    checks.append({"check": label, "passed": bool(condition)})
    if not condition:
        raise ValueError(label)


def close(a, b):
    return math.isclose(float(a), float(b), rel_tol=0, abs_tol=1e-6)


def csv_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def date_key(value):
    return value.date().isoformat() if isinstance(value, datetime) else str(value).split(" ")[0]


solution = json.loads((OUT / "question_two_solution.json").read_text(encoding="utf-8"))
check("formal causal-monthly solution", solution["diagnostics"]["parameter_mode"] == "causal_monthly")
days = {d["date"]: d for d in solution["days"]}
detail = {d: [] for d in DATES}
for row in csv_rows(OUT / "question_two_detail.csv"):
    if row["date"] in detail:
        detail[row["date"]].append(row)
daily = {r["date"]: r for r in csv_rows(OUT / "question_two_daily.csv")}
workbook = load_workbook(OUT / "result2.xlsx", read_only=True, data_only=True)
plan = {date_key(row[0]): row for row in workbook["计划购电量"].iter_rows(min_row=2, values_only=True)}
battery = {d: [] for d in DATES}
emergency = {d: [] for d in DATES}
for sheet_name, storage in (("充放电量", battery), ("紧急购电量", emergency)):
    current = None
    for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            current = date_key(row[0])
        if current in storage:
            storage[current].append(row)
workbook.close()

records = {}
for date in DATES:
    rows = sorted(detail[date], key=lambda r: int(r["period"]))
    day = days[date]
    check(f"{date} 144 ordered periods", [int(r["period"]) for r in rows] == list(range(1, 145)))
    grid = [float(r["planned_grid_kwh"]) for r in rows]
    charge = [float(r["planned_charge_kwh"]) for r in rows]
    discharge = [float(r["planned_discharge_kwh"]) for r in rows]
    check(f"{date} all planned values equal JSON and workbook", all(close(grid[t], day["grid"][t]) and close(grid[t], plan[date][t+1]) for t in range(144)))
    check(f"{date} charge and discharge equal JSON", all(close(charge[t], day["charge"][t]) and close(discharge[t], day["discharge"][t]) for t in range(144)))
    total_grid = sum(grid)
    plan_cost = sum(float(r["planned_cost_yuan"]) for r in rows)
    emergency_cost = sum(float(r["emergency_cost_yuan"]) for r in rows)
    total_emergency = sum(float(r["actual_emergency_kwh"]) for r in rows)
    check(f"{date} daily planned totals", close(total_grid, day["total_grid"]) and close(total_grid, plan[date][145]) and close(plan_cost, day["plan_cost"]) and close(plan_cost, plan[date][146]))
    check(f"{date} daily actual cost", close(plan_cost+emergency_cost, daily[date]["actual_total_cost_yuan"]) and close(total_emergency, daily[date]["actual_emergency_kwh"]))
    selected = []
    for h in HOURS:
        t = h*6
        label = f"{h}:00-{h}:10"
        check(f"{date} selected interval {label}", rows[t]["time"] == label)
        selected.append({"time": label, "grid_kwh": grid[t]})
    blocks = []
    check(f"{date} six workbook battery blocks", len(battery[date]) == 6)
    for b in range(6):
        start, end = b*24, (b+1)*24
        c, d = sum(charge[start:end]), sum(discharge[start:end])
        check(f"{date} battery block {b}", close(c, battery[date][b][2]) and close(d, battery[date][b][3]))
        blocks.append({"time": f"{b*4}:00-{(b+1)*4}:00", "charge_kwh": c, "discharge_kwh": d})
    check(f"{date} initial and terminal SOC", close(day["soc_start"], rows[0]["soc_before_kwh"]) and close(day["soc_end"], rows[-1]["soc_after_kwh"]) and close(day["soc_start"], battery[date][0][5]) and close(day["soc_end"], battery[date][1][5]))
    intervals = []
    t = 0
    while t < 144:
        if float(rows[t]["actual_emergency_kwh"]) <= 1e-7:
            t += 1
            continue
        start = t
        while t < 144 and float(rows[t]["actual_emergency_kwh"]) > 1e-7:
            t += 1
        label = rows[start]["time"].split("-")[0] + "-" + rows[t-1]["time"].split("-")[1]
        intervals.append({"time": label, "emergency_kwh": sum(float(r["actual_emergency_kwh"]) for r in rows[start:t])})
    check(f"{date} all emergency intervals equal JSON and workbook", len(intervals) == len(day["emergency_intervals"]) == len(emergency[date]) and all(r["time"] == j["period"] == x[1] and close(r["emergency_kwh"], j["amount"]) and close(r["emergency_kwh"], x[2]) for r,j,x in zip(intervals, day["emergency_intervals"], emergency[date])))
    check(f"{date} interval sums equal daily emergency", close(sum(r["emergency_kwh"] for r in intervals), total_emergency))
    records[date] = {"selected": selected, "blocks": blocks, "total_grid_kwh": total_grid, "plan_cost_yuan": plan_cost, "emergency_cost_yuan": emergency_cost, "actual_total_cost_yuan": plan_cost+emergency_cost, "total_emergency_kwh": total_emergency, "soc_start_kwh": day["soc_start"], "soc_end_kwh": day["soc_end"], "emergency_intervals": intervals}


def font(run, size=12, bold=False, east="宋体"):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), east)
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")


def para(text, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.15
    pf.keep_with_next = keep
    font(p.add_run(text), size, bold, "黑体" if bold else "宋体")
    return p


def cell_text(cell, text, size=12):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    font(p.add_run(text), size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(rows, widths, height=28):
    tbl = doc.add_table(rows=rows, cols=len(widths))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for col, width in zip(tbl.columns, widths):
        col.width = Inches(width)
    props = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:" + side)
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "8")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    props.append(borders)
    for row in tbl.rows:
        row.height = Pt(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            margins = OxmlElement("w:tcMar")
            for side, amount in (("top", 60), ("bottom", 60), ("left", 45), ("right", 45)):
                m = OxmlElement("w:"+side)
                m.set(qn("w:w"), str(amount))
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            cell._tc.get_or_add_tcPr().append(margins)
            cell_text(cell, "")
    return tbl


def f4(value):
    return f"{0.0 if abs(value) < 0.00005 else value:.4f}"


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(.55)
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.65)
for style_name in ("Normal", "Title"):
    sty = doc.styles[style_name]
    sty.font.name = "Times New Roman"
    sty.font.color.rgb = RGBColor(0, 0, 0)
    sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
doc.core_properties.title = "第二问指定日期表1表2表3"
doc.core_properties.subject = "微网计划购电 储能调度与紧急购电"
doc.core_properties.author = ""

for n, date in enumerate(DATES):
    if n:
        doc.add_page_break()
    r = records[date]
    dt = datetime.fromisoformat(date)
    para(f"{dt.year}年{dt.month}月{dt.day}日", 16, True, WD_ALIGN_PARAGRAPH.CENTER, after=18, keep=True, style="Title")
    para("表1  微网在指定时间段的购电量及全天的购电量和购电费", 12.5, True, WD_ALIGN_PARAGRAPH.CENTER, after=8, keep=True)
    para("电量单位：kWh    费用单位：元", 10, align=WD_ALIGN_PARAGRAPH.RIGHT, after=5, keep=True)
    t1 = table(4, [7.4/6]*6, 32)
    for i, text in enumerate(["时间段", "购电量"]*3):
        cell_text(t1.cell(0,i), text)
    for j, entry in enumerate(r["selected"]):
        row, col = 1+j//3, 2*(j%3)
        cell_text(t1.cell(row,col), entry["time"])
        cell_text(t1.cell(row,col+1), f4(entry["grid_kwh"]))
    cell_text(t1.cell(3,0).merge(t1.cell(3,1)), "全天购电量")
    cell_text(t1.cell(3,2), f4(r["total_grid_kwh"]))
    cell_text(t1.cell(3,3).merge(t1.cell(3,4)), "全天购电费")
    cell_text(t1.cell(3,5), f"{r['plan_cost_yuan']:.2f}")
    para("注：表1统计正常计划购电量及其费用，紧急购电量单列于表3。", 10, before=7, after=21)
    para("表2  储能设备在指定时间段的充放电量及0:00和24:00的储电量", 12.5, True, WD_ALIGN_PARAGRAPH.CENTER, after=8, keep=True)
    para("电量单位：kWh", 10, align=WD_ALIGN_PARAGRAPH.RIGHT, after=5, keep=True)
    t2 = table(5, [7.4/6]*6, 32)
    for i, text in enumerate(["时间段", "充电量", "放电量"]*2):
        cell_text(t2.cell(0,i), text)
    for j, entry in enumerate(r["blocks"]):
        row, col = 1+j//2, 3*(j%2)
        cell_text(t2.cell(row,col), entry["time"])
        cell_text(t2.cell(row,col+1), f4(entry["charge_kwh"]))
        cell_text(t2.cell(row,col+2), f4(entry["discharge_kwh"]))
    cell_text(t2.cell(4,0).merge(t2.cell(4,1)), "0:00 储电量")
    cell_text(t2.cell(4,2), f4(r["soc_start_kwh"]))
    cell_text(t2.cell(4,3).merge(t2.cell(4,4)), "24:00 储电量")
    cell_text(t2.cell(4,5), f4(r["soc_end_kwh"]))
    para(f"当日紧急购电费为{r['emergency_cost_yuan']:.2f}元；含紧急购电的实际总费用为{r['actual_total_cost_yuan']:.2f}元。", 10, before=10, after=4)
    para("注：电量保留四位小数，费用保留两位小数；汇总按未舍入原值计算。", 10, after=0)

doc.add_page_break()
para("表3  微网在指定日期的紧急购电量", 15, True, WD_ALIGN_PARAGRAPH.CENTER, after=14, keep=True, style="Title")
para("电量单位：kWh", 10, align=WD_ALIGN_PARAGRAPH.RIGHT, after=5, keep=True)
max_rows = max(len(r["emergency_intervals"]) for r in records.values())
t3 = table(2+max_rows, [1.0,.85]*4, 24)
for col, date in enumerate(DATES):
    dt = datetime.fromisoformat(date)
    cell_text(t3.cell(0,2*col).merge(t3.cell(0,2*col+1)), f"{dt.year}.{dt.month}.{dt.day}", 12.5)
    cell_text(t3.cell(1,2*col), "时间段", 11.5)
    cell_text(t3.cell(1,2*col+1), "购电量", 11.5)
    for j, entry in enumerate(records[date]["emergency_intervals"]):
        cell_text(t3.cell(j+2,2*col), entry["time"], 11)
        cell_text(t3.cell(j+2,2*col+1), f4(entry["emergency_kwh"]), 11)
for row in t3.rows[:2]:
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
para("注：连续发生的10分钟紧急购电时段合并列示；空白单元格表示该日期无更多紧急购电记录。", 10, before=9, after=5)
para("四日紧急购电量合计依次为："+"、".join(f4(records[d]["total_emergency_kwh"]) for d in DATES)+" kWh。", 10, after=0)
DEST.parent.mkdir(parents=True, exist_ok=True)
# The bundled default template carries a decorative Title border and some
# conditional cell-border rules. Override both to match the supplied plain grid.
for element in (doc.styles.element, doc.element):
    for border in list(element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
for tbl in doc.tables:
    for tc in tbl._tbl.xpath(".//w:tc"):
        props = tc.get_or_add_tcPr()
        for old_border in list(props.findall(qn("w:tcBorders"))):
            props.remove(old_border)
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement("w:"+side)
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "8")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        props.append(borders)
doc.save(DEST)
check("nine editable Word tables", len(Document(DEST).tables) == 9)
check("table three complete row count", len(Document(DEST).tables[-1].rows) == 19)
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
audit = {"status": "PASS", "table_1_count": 4, "table_2_count": 4, "table_3_count": 1, "table_3_interval_counts": {d:len(records[d]["emergency_intervals"]) for d in DATES}, "checks":checks, "source_sha256":{p.name:sha(p) for p in (OUT/"question_two_solution.json", OUT/"question_two_daily.csv", OUT/"question_two_detail.csv", OUT/"result2.xlsx")}, "document_sha256":sha(DEST), "rounding":"4 decimals for kWh; 2 decimals for yuan; totals calculated before rounding", "table_1_cost_basis":"planned normal grid purchase cost; emergency and actual total cost separately stated in notes"}
(ROOT/"checks/paper_tables_qa/data_verification.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding="utf-8")
(OUT/"paper_tables/paper_tables_1_2_3_data.json").write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"document":str(DEST),"checks_passed":len(checks),"table_3_interval_counts":audit["table_3_interval_counts"]},ensure_ascii=False))
