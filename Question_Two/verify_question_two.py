from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser(description="核验C题第二问result2.xlsx")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--solution", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    solution = json.loads(args.solution.read_text(encoding="utf-8"))
    days = solution["days"]
    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    required = ["计划购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != required:
        raise AssertionError(f"工作表名称或顺序错误: {workbook.sheetnames}")

    plan = workbook["计划购电量"]
    battery = workbook["充放电量"]
    emergency = workbook["紧急购电量"]
    plan.calculate_dimension(force=True)
    battery.calculate_dimension(force=True)
    emergency.calculate_dimension(force=True)
    if plan.max_row != 335 or plan.max_column != 147:
        raise AssertionError(f"计划购电量尺寸错误: {plan.max_row}×{plan.max_column}")
    if battery.max_row < 2005 or battery.max_column != 6:
        raise AssertionError(f"充放电量尺寸错误: {battery.max_row}×{battery.max_column}")

    max_grid_error = 0.0
    max_total_error = 0.0
    plan_values = plan.iter_rows(
        min_row=2, max_row=335, min_col=1, max_col=147, values_only=True
    )
    for values, day in zip(plan_values, days, strict=True):
        for actual_value, expected in zip(values[1:145], day["grid"], strict=True):
            actual = float(actual_value)
            max_grid_error = max(max_grid_error, abs(actual - expected))
        max_total_error = max(
            max_total_error,
            abs(float(values[145]) - float(day["total_grid"])),
            abs(float(values[146]) - float(day["plan_cost"])),
        )

    max_block_error = 0.0
    battery_values = battery.iter_rows(
        min_row=2, max_row=2005, min_col=1, max_col=6, values_only=True
    )
    for day in days:
        for block in range(6):
            values = next(battery_values)
            expected_charge = sum(day["charge"][block * 24 : (block + 1) * 24])
            expected_discharge = sum(day["discharge"][block * 24 : (block + 1) * 24])
            max_block_error = max(
                max_block_error,
                abs(float(values[2]) - expected_charge),
                abs(float(values[3]) - expected_discharge),
            )

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    raise AssertionError(f"发现Excel错误值: {sheet.title}!{cell.coordinate}={cell.value}")
                if isinstance(cell.value, float) and not math.isfinite(cell.value):
                    raise AssertionError(f"发现非有限数值: {sheet.title}!{cell.coordinate}")
    plan_rows = plan.max_row
    battery_rows = battery.max_row
    emergency_rows = emergency.max_row
    workbook.close()

    checks = {
        "status": "PASS",
        "workbook": str(args.workbook.resolve()),
        "formal_days": len(days),
        "plan_rows": plan_rows,
        "battery_rows": battery_rows,
        "emergency_rows": emergency_rows,
        "max_grid_error_kwh": max_grid_error,
        "max_daily_total_error": max_total_error,
        "max_block_error_kwh": max_block_error,
        "model_max_balance_residual_kwh": solution["diagnostics"]["max_balance_residual_kwh"],
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
