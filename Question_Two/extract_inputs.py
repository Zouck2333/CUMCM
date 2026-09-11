from __future__ import annotations

import argparse
import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


DELTA_H = 1.0 / 6.0


def _number(value: object, *, context: str) -> float:
    if value is None:
        raise ValueError(f"缺少数值: {context}")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"无法转换为数值: {context}={value!r}") from exc


def _date_text(value: object, *, context: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return datetime.fromisoformat(value).date().isoformat()
    raise ValueError(f"无法识别日期: {context}={value!r}")


def _time_text(value: object) -> str:
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, datetime):
        return value.time().strftime("%H:%M:%S")
    return str(value)


def load_attachment1(path: Path) -> dict[str, list[float] | list[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    if sheet.max_row < 145 or sheet.max_column < 4:
        raise ValueError("附件1必须包含表头和144个时段、4列数据")

    times: list[str] = []
    price: list[float] = []
    prior_load: list[float] = []
    prior_pv: list[float] = []
    for row, values in enumerate(
        sheet.iter_rows(min_row=2, max_row=145, min_col=1, max_col=4, values_only=True),
        start=2,
    ):
        times.append(_time_text(values[0]))
        price.append(_number(values[1], context=f"附件1 B{row}"))
        prior_load.append(
            _number(values[2], context=f"附件1 C{row}")
            * DELTA_H
        )
        prior_pv.append(
            max(
                0.0,
                _number(values[3], context=f"附件1 D{row}")
                * DELTA_H,
            )
        )
    workbook.close()
    return {
        "times": times,
        "price": price,
        "prior_load": prior_load,
        "prior_pv": prior_pv,
    }


def load_attachment2(path: Path) -> tuple[list[str], list[list[float]], list[list[float]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if len(workbook.sheetnames) < 2:
        raise ValueError("附件2必须包含负载和光伏两个工作表")

    load_sheet = workbook[workbook.sheetnames[0]]
    pv_sheet = workbook[workbook.sheetnames[1]]
    if load_sheet.max_row < 366 or pv_sheet.max_row < 366:
        raise ValueError("附件2必须包含2025年365天数据")
    if load_sheet.max_column < 145 or pv_sheet.max_column < 145:
        raise ValueError("附件2每天必须包含144个时段")

    dates: list[str] = []
    load: list[list[float]] = []
    pv: list[list[float]] = []
    load_rows = load_sheet.iter_rows(
        min_row=2, max_row=366, min_col=1, max_col=145, values_only=True
    )
    pv_rows = pv_sheet.iter_rows(
        min_row=2, max_row=366, min_col=1, max_col=145, values_only=True
    )
    for row, (load_values, pv_values) in enumerate(
        zip(load_rows, pv_rows, strict=True), start=2
    ):
        load_date = _date_text(load_values[0], context=f"负载 A{row}")
        pv_date = _date_text(pv_values[0], context=f"光伏 A{row}")
        if load_date != pv_date:
            raise ValueError(f"附件2日期不一致: row={row}, {load_date} != {pv_date}")
        dates.append(load_date)
        load.append(
            [
                _number(load_values[col - 1], context=f"负载 row={row}, col={col}")
                * DELTA_H
                for col in range(2, 146)
            ]
        )
        pv.append(
            [
                max(
                    0.0,
                    _number(pv_values[col - 1], context=f"光伏 row={row}, col={col}")
                    * DELTA_H,
                )
                for col in range(2, 146)
            ]
        )
    workbook.close()

    expected_start = date(2025, 1, 1).isoformat()
    expected_end = date(2025, 12, 31).isoformat()
    if dates[0] != expected_start or dates[-1] != expected_end:
        raise ValueError(f"附件2日期范围错误: {dates[0]} 至 {dates[-1]}")
    return dates, load, pv


def load_template_headers(path: Path) -> list[str]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [
        str(value)
        for value in next(
            sheet.iter_rows(min_row=1, max_row=1, min_col=2, max_col=145, values_only=True)
        )
    ]
    workbook.close()
    if len(headers) != 144 or any(value in {"None", ""} for value in headers):
        raise ValueError("result2模板计划购电量工作表缺少144个时段标题")
    return headers


def main() -> None:
    parser = argparse.ArgumentParser(description="提取C题第二问输入数据")
    parser.add_argument("--attachment1", type=Path, required=True)
    parser.add_argument("--attachment2", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    attachment1 = load_attachment1(args.attachment1)
    dates, load, pv = load_attachment2(args.attachment2)
    payload = {
        "delta_h": DELTA_H,
        "dates": dates,
        "input_times": attachment1["times"],
        "time_labels": load_template_headers(args.template),
        "price": attachment1["price"],
        "prior_load": attachment1["prior_load"],
        "prior_pv": attachment1["prior_pv"],
        "actual_load": load,
        "actual_pv": pv,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(f"已提取输入数据: {args.output}")


if __name__ == "__main__":
    main()
