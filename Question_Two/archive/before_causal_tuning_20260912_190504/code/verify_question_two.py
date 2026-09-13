from __future__ import annotations

import argparse
import csv
import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import numpy as np


N_TIME = 144
FORMAL_START = date(2025, 2, 1)
FORMAL_END = date(2025, 12, 31)
REQUIRED_SHEETS = ["计划购电量", "充放电量", "紧急购电量"]
BLOCK_LABELS = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]
DETAIL_FIELDS = {
    "date",
    "period",
    "time",
    "price_yuan_per_kwh",
    "predicted_load_kwh",
    "actual_load_kwh",
    "predicted_pv_kwh",
    "actual_pv_kwh",
    "planned_grid_kwh",
    "planned_charge_kwh",
    "planned_discharge_kwh",
    "predicted_curtailment_kwh",
    "soc_before_kwh",
    "soc_after_kwh",
    "actual_emergency_kwh",
    "actual_surplus_kwh",
    "planned_cost_yuan",
    "emergency_cost_yuan",
}
DAILY_FIELDS = {
    "date",
    "warmup",
    "soc_start_kwh",
    "soc_end_kwh",
    "planned_grid_kwh",
    "actual_emergency_kwh",
    "actual_surplus_kwh",
    "planned_cost_yuan",
    "emergency_cost_yuan",
    "actual_total_cost_yuan",
    "reserve_kwh",
    "solver_status",
}
ERROR_TEXT = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
}
TIME_LABEL_PATTERN = re.compile(
    r"^\s*(\d{1,2}):(\d{1,2})(\+1)?\s*-\s*"
    r"(\d{1,2}):(\d{1,2})(\+1)?\s*$"
)


class VerificationError(RuntimeError):
    """输入或结果结构不足以继续核验。"""


@dataclass
class Audit:
    checks: dict[str, dict[str, Any]] = field(default_factory=dict)
    metrics: dict[str, Any] = field(default_factory=dict)
    failures: list[str] = field(default_factory=list)

    def check(self, name: str, passed: bool, detail: str) -> None:
        status = "PASS" if passed else "FAIL"
        self.checks[name] = {"status": status, "detail": detail}
        if not passed:
            self.failures.append(f"{name}: {detail}")


@dataclass
class Maximum:
    value: float = 0.0
    location: str | None = None

    def add(self, value: float, location: str) -> None:
        absolute = abs(float(value))
        if absolute > self.value:
            self.value = absolute
            self.location = location


def _must(condition: bool, message: str) -> None:
    if not condition:
        raise VerificationError(message)


def _read_json(path: Path) -> dict[str, Any]:
    _must(path.is_file(), f"文件不存在: {path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise VerificationError(f"无法读取JSON文件 {path}: {exc}") from exc
    _must(isinstance(payload, dict), f"JSON根节点必须是对象: {path}")
    return payload


def _read_csv(path: Path, required_fields: set[str]) -> list[dict[str, str]]:
    _must(path.is_file(), f"文件不存在: {path}")
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            fields = set(reader.fieldnames or [])
            missing = sorted(required_fields - fields)
            _must(not missing, f"{path.name}缺少列: {missing}")
            return list(reader)
    except OSError as exc:
        raise VerificationError(f"无法读取CSV文件 {path}: {exc}") from exc


def _number(value: Any, context: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise VerificationError(f"{context}不是数值: {value!r}") from exc
    if not math.isfinite(result):
        raise VerificationError(f"{context}不是有限数值: {value!r}")
    return result


def _iso_date(value: Any, context: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        text = value.strip()
        try:
            return datetime.fromisoformat(text).date().isoformat()
        except ValueError as exc:
            raise VerificationError(f"{context}无法识别为日期: {value!r}") from exc
    raise VerificationError(f"{context}无法识别为日期: {value!r}")


def _continuous_dates(values: list[str]) -> bool:
    parsed = [date.fromisoformat(value) for value in values]
    return all(
        right - left == timedelta(days=1)
        for left, right in zip(parsed, parsed[1:])
    )


def _parse_clock(hour_text: str, minute_text: str, plus_one: str | None) -> int:
    hour = int(hour_text)
    minute = int(minute_text)
    _must(0 <= hour <= 24, f"小时越界: {hour_text}:{minute_text}")
    _must(0 <= minute < 60, f"分钟越界: {hour_text}:{minute_text}")
    _must(
        hour < 24 or minute == 0,
        f"24点只能写作24:00: {hour_text}:{minute_text}",
    )
    value = hour * 60 + minute
    if plus_one:
        value += 24 * 60
    return value


def _validate_time_labels(labels: list[str]) -> tuple[bool, str]:
    if len(labels) != N_TIME:
        return False, f"时段标题数量应为{N_TIME}，实际为{len(labels)}"
    previous_end: int | None = None
    first_start: int | None = None
    for index, label in enumerate(labels, start=1):
        match = TIME_LABEL_PATTERN.match(label)
        if match is None:
            return False, f"第{index}个时段标题格式错误: {label!r}"
        try:
            start = _parse_clock(match[1], match[2], match[3])
            end = _parse_clock(match[4], match[5], match[6])
        except VerificationError as exc:
            return False, f"第{index}个时段标题错误: {exc}"
        if previous_end is not None:
            while start < previous_end:
                start += 24 * 60
            if start != previous_end:
                return (
                    False,
                    f"第{index - 1}与第{index}个时段不连续: "
                    f"{labels[index - 2]!r}, {label!r}",
                )
        while end <= start:
            end += 24 * 60
        if end - start != 10:
            return False, f"第{index}个时段长度不是10分钟: {label!r}"
        if first_start is None:
            first_start = start
        previous_end = end
    assert first_start is not None and previous_end is not None
    if previous_end - first_start != 24 * 60:
        return (
            False,
            f"144个时段总跨度不是24小时: {previous_end - first_start}分钟",
        )
    if len(set(labels)) != N_TIME:
        return False, "144个时段标题存在重复"
    return True, f"{N_TIME}个标题连续覆盖24小时，每段10分钟"


def _merge_emergency(
    emergency: list[float], time_labels: list[str], threshold: float
) -> list[dict[str, float | str]]:
    intervals: list[dict[str, float | str]] = []
    start: int | None = None
    for index, value in enumerate(emergency):
        active = value > threshold
        if active and start is None:
            start = index
        if start is not None and (not active or index == len(emergency) - 1):
            end = index if active and index == len(emergency) - 1 else index - 1
            first = time_labels[start].split("-", 1)[0]
            last = time_labels[end].split("-", 1)[-1]
            intervals.append(
                {
                    "period": f"{first}-{last}",
                    "amount": float(sum(emergency[start : end + 1])),
                }
            )
            start = None
    return intervals


def _resolve_companion(
    explicit: Path | None, parent: Path, filename: str
) -> Path:
    return explicit if explicit is not None else parent / filename


def verify_purchase_risk(audit, solution, inputs, daily_by_date, detail_by_date, abs_tol):
    """Independent nearest-rank implementation; do not import the solver."""
    diag = solution["diagnostics"]
    if diag.get("purchase_strategy", "legacy_scenarios") == "legacy_scenarios":
        excess = max(
            (float(row["predicted_curtailment_kwh"]) - float(row["predicted_pv_kwh"])
             for rows in detail_by_date.values() for row in rows), default=0.0)
        audit.check("legacy_pv_surplus_upper_bound", excess <= abs_tol, f"最大违反={excess}")
        return
    _must(diag.get("predicted_surplus_mode") == "total_supply", "校准策略必须使用总供能富余口径")
    history = solution.get("forecast_history", [])
    _must([row["date"] for row in history] == inputs["dates"], "预测历史必须覆盖365天")
    for row in history:
        for field_name in ("predicted_load", "predicted_pv"):
            _must(len(row[field_name]) == N_TIME, "历史预测长度错误")
            _must(all(_number(v, field_name) >= 0 for v in row[field_name]), "历史预测负值")
    dates = [date.fromisoformat(text) for text in inputs["dates"]]
    errors = [[inputs["actual_load"][k][t] - inputs["actual_pv"][k][t]
               - history[k]["predicted_load"][t] + history[k]["predicted_pv"][t]
               for t in range(N_TIME)] for k in range(365)]
    quantile = float(diag["purchase_quantile"])
    window = int(diag["risk_window_days"])
    radius = int(diag["risk_radius_periods"])
    _must(0 < quantile < 1 and window >= 1 and 0 <= radius < N_TIME, "风险参数不合法")
    _must(diag["purchase_quantile_method"] == "inverted_cdf", "风险分位数算法不一致")
    margin_error, risk_violation, forecast_error, objective_error = Maximum(), Maximum(), Maximum(), Maximum()
    source_errors = 0
    for day in solution["days"]:
        target = inputs["dates"].index(day["date"])
        candidates = list(range(max(1, target - window), target))
        same = [k for k in candidates if (dates[k].weekday() < 5) == (dates[target].weekday() < 5)]
        grouping = diag.get("risk_grouping", "legacy")
        _must(grouping in {"all", "legacy"}, "残差分组参数错误")
        selected = same if grouping == "legacy" and len(same) >= 5 else candidates
        daily = daily_by_date[day["date"]]
        sources = [inputs["dates"][k] for k in selected]
        if (daily["risk_source_dates"].split(";") != sources
                or int(daily["risk_source_day_count"]) != len(selected)
                or daily["purchase_strategy"] != "calibrated_quantile"
                or int(daily["scenario_count"]) != 0):
            source_errors += 1
        _must(len(day["purchase_margin"]) == N_TIME, "JSON风险余量长度错误")
        margin_sum = 0.0
        for t, row in enumerate(detail_by_date[day["date"]]):
            sample = sorted(errors[k][u] for k in selected
                            for u in range(max(0, t-radius), min(N_TIME, t+radius+1)))
            expected = max(0.0, sample[max(0, math.ceil(quantile * len(sample)) - 1)])
            actual = _number(row["purchase_margin_kwh"], "风险余量")
            margin_sum += actual
            margin_error.add(actual - expected, f"{day['date']}#{t+1}")
            margin_error.add(float(day["purchase_margin"][t]) - actual, "JSON风险余量")
            risk_violation.add(max(0.0, expected - float(row["predicted_curtailment_kwh"])), "风险下限")
            forecast_error.add(float(row["predicted_load_kwh"]) - history[target]["predicted_load"][t], "负荷预测历史")
            forecast_error.add(float(row["predicted_pv_kwh"]) - history[target]["predicted_pv"][t], "光伏预测历史")
        margin_error.add(margin_sum - float(daily["purchase_margin_kwh"]), "每日风险余量")
        objective_error.add(float(daily["solver_objective_yuan"]) - float(daily["planned_cost_yuan"]), "校准策略主费用")
    audit.check("purchase_risk_sources", source_errors == 0, f"风险样本来源错误天数={source_errors}")
    audit.check("purchase_margin_recomputed", margin_error.value <= abs_tol, f"独立重算最大误差={margin_error.value}")
    audit.check("calibrated_supply_floor", risk_violation.value <= abs_tol, f"供能余量最大不足={risk_violation.value}")
    audit.check("forecast_history_mapping", forecast_error.value <= abs_tol, f"预测历史最大误差={forecast_error.value}")
    audit.check("calibrated_primary_objective", objective_error.value <= abs_tol, f"费用最大误差={objective_error.value}")
    if diag.get("forecast_method") == "calendar_trend":
        # Independent augmented least squares, rather than the solver's normal equations.
        load = np.asarray(inputs["actual_load"], dtype=float)
        pv = np.asarray(inputs["actual_pv"], dtype=float)
        load_error, pv_error = Maximum(), Maximum()
        window, degree, pv_window = (int(diag[k]) for k in
                                    ("load_window_days", "load_trend_degree", "pv_window_days"))
        training_source_errors = 0
        for target in range(31, 365):
            indices = list(range(max(0, target-window), target))
            tau = np.asarray([(k-target)/28.0 for k in indices])
            x = np.column_stack([np.ones(len(indices))] + [tau ** j for j in range(1,degree+1)]
                                + [[float(dates[k].weekday() == j) for k in indices] for j in range(6)])
            ridge = np.eye(x.shape[1]) * .1
            ridge[0,0] = .0001
            coefficients = np.linalg.lstsq(np.vstack((x,ridge)),
                np.vstack((np.log(np.maximum(1,load[indices])), np.zeros((x.shape[1],144)))), rcond=None)[0]
            query = np.asarray([1.] + [0.]*degree + [float(dates[target].weekday()==j) for j in range(6)])
            expected_load = np.exp(query @ coefficients)
            pv_indices = list(range(max(0,target-pv_window),target))
            z = np.column_stack((np.ones(len(pv_indices)), (np.asarray(pv_indices)-target)/pv_window))
            pv_coef = np.linalg.lstsq(np.vstack((z,np.diag([.0001,math.sqrt(.3)]))),
                np.vstack((pv[pv_indices],np.zeros((2,144)))), rcond=None)[0]
            expected_pv = np.maximum(0,pv_coef[0])
            load_error.add(float(np.max(np.abs(expected_load-history[target]["predicted_load"]))), inputs["dates"][target])
            pv_error.add(float(np.max(np.abs(expected_pv-history[target]["predicted_pv"]))), inputs["dates"][target])
            row = daily_by_date[inputs["dates"][target]]
            if (row["similar_load_dates"].split(";") != [inputs["dates"][k] for k in indices]
                or row["similar_pv_dates"].split(";") != [inputs["dates"][k] for k in pv_indices]
                or row["forecast_method"] != "calendar_trend"):
                training_source_errors += 1
        audit.check("calendar_load_forecast_recomputed", load_error.value <= abs_tol, f"独立最小二乘重算误差={load_error.value}")
        audit.check("recent_pv_forecast_recomputed", pv_error.value <= abs_tol, f"独立最小二乘重算误差={pv_error.value}")
        audit.check("forecast_training_sources", training_source_errors == 0, f"训练来源错误天数={training_source_errors}")


def verify(
    *,
    workbook_path: Path,
    solution_path: Path,
    input_path: Path,
    detail_path: Path,
    daily_path: Path,
    abs_tol: float,
    interval_eps: float,
) -> dict[str, Any]:
    audit = Audit()
    solution = _read_json(solution_path)
    inputs = _read_json(input_path)
    detail_rows = _read_csv(detail_path, DETAIL_FIELDS)
    daily_rows = _read_csv(daily_path, DAILY_FIELDS)

    days = solution.get("days")
    diagnostics = solution.get("diagnostics")
    _must(isinstance(days, list), "solution.days必须是数组")
    _must(isinstance(diagnostics, dict), "solution.diagnostics必须是对象")

    input_dates = [str(value) for value in inputs.get("dates", [])]
    _must(
        len(input_dates) == 365,
        f"input_data.json应包含365个日期，实际为{len(input_dates)}",
    )
    input_expected = [
        (date(2025, 1, 1) + timedelta(days=index)).isoformat()
        for index in range(365)
    ]
    audit.check(
        "input_dates_continuous_unique",
        input_dates == input_expected and len(set(input_dates)) == len(input_dates),
        f"日期范围={input_dates[0]}至{input_dates[-1]}，数量={len(input_dates)}",
    )

    time_labels = [str(value) for value in inputs.get("time_labels", [])]
    time_ok, time_detail = _validate_time_labels(time_labels)
    audit.check("time_headers_144_continuous", time_ok, time_detail)
    _must(len(time_labels) == N_TIME, "缺少144个标准时段标题")

    price = [
        _number(value, f"input.price[{index}]")
        for index, value in enumerate(inputs.get("price", []))
    ]
    _must(len(price) == N_TIME, f"电价数量应为144，实际为{len(price)}")
    delta_h = _number(inputs.get("delta_h"), "input.delta_h")
    eta_c = _number(diagnostics.get("eta_charge"), "diagnostics.eta_charge")
    eta_d = _number(diagnostics.get("eta_discharge"), "diagnostics.eta_discharge")
    soc_min = _number(
        diagnostics.get("soc_min_kwh"), "diagnostics.soc_min_kwh"
    )
    soc_max = _number(
        diagnostics.get("soc_max_kwh"), "diagnostics.soc_max_kwh"
    )
    energy_limit = 5000.0 * delta_h

    actual_load_input = inputs.get("actual_load")
    actual_pv_input = inputs.get("actual_pv")
    _must(
        isinstance(actual_load_input, list) and len(actual_load_input) == 365,
        "input.actual_load维度必须为365×144",
    )
    _must(
        isinstance(actual_pv_input, list) and len(actual_pv_input) == 365,
        "input.actual_pv维度必须为365×144",
    )
    _must(
        all(
            isinstance(row, list) and len(row) == N_TIME
            for row in actual_load_input
        ),
        "input.actual_load维度必须为365×144",
    )
    _must(
        all(
            isinstance(row, list) and len(row) == N_TIME
            for row in actual_pv_input
        ),
        "input.actual_pv维度必须为365×144",
    )

    formal_expected = [
        value
        for value in input_dates
        if FORMAL_START.isoformat() <= value <= FORMAL_END.isoformat()
    ]
    solution_dates = [str(day.get("date")) for day in days]
    audit.check(
        "solution_dates_continuous_unique",
        solution_dates == formal_expected
        and len(set(solution_dates)) == len(solution_dates)
        and _continuous_dates(solution_dates),
        f"日期范围={solution_dates[0] if solution_dates else None}至"
        f"{solution_dates[-1] if solution_dates else None}，数量={len(solution_dates)}",
    )
    _must(
        solution_dates == formal_expected,
        "solution日期不完整或顺序错误，无法进行Excel逐行映射核验",
    )

    detail_by_date: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in detail_rows:
        detail_by_date[row["date"]].append(row)
    detail_dates = list(detail_by_date)
    audit.check(
        "detail_dates_and_row_count",
        detail_dates == formal_expected
        and len(detail_rows) == len(formal_expected) * N_TIME
        and all(len(detail_by_date[value]) == N_TIME for value in formal_expected),
        f"日期数={len(detail_dates)}，逐时行数={len(detail_rows)}",
    )
    _must(
        all(
            value in detail_by_date and len(detail_by_date[value]) == N_TIME
            for value in formal_expected
        ),
        "逐时CSV日期或行数不完整",
    )

    daily_by_date = {row["date"]: row for row in daily_rows}
    audit.check(
        "daily_dates_continuous_unique",
        [row["date"] for row in daily_rows] == input_dates
        and len(daily_by_date) == len(daily_rows),
        f"每日CSV日期数={len(daily_rows)}，唯一日期数={len(daily_by_date)}",
    )
    _must(
        all(value in daily_by_date for value in formal_expected),
        "每日CSV缺少正式评价日期",
    )

    daily_all_soc_link_error = Maximum()
    daily_endpoint_bound_violation = Maximum()
    solver_status_failures = 0
    previous_daily_soc_end: float | None = None
    for day_index, row in enumerate(daily_rows):
        day_text = row["date"]
        daily_soc_start = _number(
            row["soc_start_kwh"], f"daily {day_text}.soc_start"
        )
        daily_soc_end = _number(
            row["soc_end_kwh"], f"daily {day_text}.soc_end"
        )
        expected_start = 6000.0 if day_index == 0 else previous_daily_soc_end
        assert expected_start is not None
        daily_all_soc_link_error.add(
            daily_soc_start - expected_start, f"daily {day_text}日初"
        )
        daily_endpoint_bound_violation.add(
            max(
                0.0,
                soc_min - daily_soc_start,
                soc_min - daily_soc_end,
                daily_soc_start - soc_max,
                daily_soc_end - soc_max,
            ),
            f"daily {day_text}",
        )
        previous_daily_soc_end = daily_soc_end
        try:
            if int(row["solver_status"]) != 0:
                solver_status_failures += 1
        except (TypeError, ValueError):
            solver_status_failures += 1
    audit.check(
        "daily_soc_cross_day_link_all_year",
        daily_all_soc_link_error.value <= abs_tol,
        f"最大误差={daily_all_soc_link_error.value:.12g} kWh，"
        f"位置={daily_all_soc_link_error.location}",
    )
    audit.check(
        "daily_soc_endpoint_bounds_all_year",
        daily_endpoint_bound_violation.value <= abs_tol,
        f"最大违反={daily_endpoint_bound_violation.value:.12g} kWh，"
        f"位置={daily_endpoint_bound_violation.location}",
    )

    # 新版求解器会保存每日历史来源。存在这些列时，核验器同时执行
    # 信息边界检查；旧版结果没有这些列时仍可完成其余物理核验。
    history_fields = {
        "similar_load_dates",
        "similar_pv_dates",
        "scenario_source_dates",
        "reserve_latest_source_date",
        "scenario_count",
    }
    available_daily_fields = set(daily_rows[0]) if daily_rows else set()
    if history_fields <= available_daily_fields:
        history_boundary_violations = 0
        history_count_violations = 0
        reserve_source_violations = 0
        for day_index, row in enumerate(daily_rows):
            target = date.fromisoformat(row["date"])
            for field_name in (
                "similar_load_dates",
                "similar_pv_dates",
                "scenario_source_dates",
            ):
                sources = [
                    value
                    for value in row[field_name].split(";")
                    if value.strip()
                ]
                for source in sources:
                    if date.fromisoformat(source) >= target:
                        history_boundary_violations += 1
                if field_name == "scenario_source_dates":
                    scenario_count = int(row["scenario_count"])
                    # 1月1日没有历史残差，程序使用一个确定性回退情景。
                    expected_count = 1 if day_index == 0 else len(sources)
                    if scenario_count != expected_count:
                        history_count_violations += 1
            reserve_source = row["reserve_latest_source_date"].strip()
            expected_reserve_source = (
                "" if day_index == 0 else input_dates[day_index - 1]
            )
            if reserve_source != expected_reserve_source:
                reserve_source_violations += 1
            if reserve_source and date.fromisoformat(reserve_source) >= target:
                history_boundary_violations += 1
        audit.check(
            "history_information_boundary",
            history_boundary_violations == 0
            and history_count_violations == 0
            and reserve_source_violations == 0,
            f"未来来源={history_boundary_violations}，情景数不一致="
            f"{history_count_violations}，储备最新来源错误="
            f"{reserve_source_violations}",
        )
    else:
        missing_history_fields = sorted(history_fields - available_daily_fields)
        audit.check(
            "history_information_boundary",
            True,
            f"旧版daily未提供{missing_history_fields}，本项跳过",
        )

    lexicographic_fields = {
        "storage_throughput_kwh",
        "peak_grid_kwh",
        "objective_mode",
    }
    check_lexicographic_metrics = lexicographic_fields <= available_daily_fields
    verify_purchase_risk(audit, solution, inputs, daily_by_date, detail_by_date, abs_tol)

    period_mismatches = 0
    time_mismatches = 0
    finite_failures = 0
    negative_failures = 0
    mutual_exclusion_count = 0
    complementarity_count = 0
    throughput_error = Maximum()
    peak_grid_error = Maximum()
    objective_mode_mismatches = 0
    terminal_reserve_violation = Maximum()
    power_limit_excess = Maximum()
    soc_lower_violation = Maximum()
    soc_upper_violation = Maximum()
    soc_recursion_error = Maximum()
    soc_within_day_link_error = Maximum()
    soc_cross_day_error = Maximum()
    predicted_balance_error = Maximum()
    actual_balance_error = Maximum()
    plan_cost_error = Maximum()
    emergency_cost_error = Maximum()
    input_price_error = Maximum()
    input_actual_load_error = Maximum()
    input_actual_pv_error = Maximum()
    json_grid_error = Maximum()
    json_charge_error = Maximum()
    json_discharge_error = Maximum()
    json_soc_endpoint_error = Maximum()
    json_total_grid_error = Maximum()
    json_plan_cost_error = Maximum()
    interval_label_mismatches = 0
    interval_amount_error = Maximum()
    daily_grid_error = Maximum()
    daily_emergency_error = Maximum()
    daily_surplus_error = Maximum()
    daily_plan_cost_error = Maximum()
    daily_emergency_cost_error = Maximum()
    daily_total_cost_error = Maximum()
    aggregate_by_date: dict[str, dict[str, float]] = {}
    reconstructed_intervals: dict[
        str, list[dict[str, float | str]]
    ] = {}

    previous_soc_end: float | None = None
    total_grid = 0.0
    total_emergency = 0.0
    total_plan_cost = 0.0
    total_emergency_cost = 0.0
    numeric_fields = DETAIL_FIELDS - {"date", "period", "time"}

    for day, day_text in zip(days, formal_expected, strict=True):
        rows = detail_by_date[day_text]
        grid = day.get("grid")
        charge = day.get("charge")
        discharge = day.get("discharge")
        _must(
            isinstance(grid, list) and len(grid) == N_TIME,
            f"{day_text} JSON grid长度不是144",
        )
        _must(
            isinstance(charge, list) and len(charge) == N_TIME,
            f"{day_text} JSON charge长度不是144",
        )
        _must(
            isinstance(discharge, list) and len(discharge) == N_TIME,
            f"{day_text} JSON discharge长度不是144",
        )
        soc_start_json = _number(day.get("soc_start"), f"{day_text}.soc_start")
        soc_end_json = _number(day.get("soc_end"), f"{day_text}.soc_end")
        if previous_soc_end is not None:
            soc_cross_day_error.add(
                soc_start_json - previous_soc_end, f"{day_text}日初"
            )
        previous_soc_end = soc_end_json

        date_grid = 0.0
        date_emergency = 0.0
        date_surplus = 0.0
        date_plan_cost = 0.0
        date_emergency_cost = 0.0
        emergency_values: list[float] = []
        previous_soc_after: float | None = None
        source_index = input_dates.index(day_text)

        for period_index, row in enumerate(rows):
            location = f"{day_text}#{period_index + 1}"
            try:
                period_value = int(row["period"])
            except (TypeError, ValueError):
                period_value = -1
            if period_value != period_index + 1:
                period_mismatches += 1
            if row["time"] != time_labels[period_index]:
                time_mismatches += 1

            values: dict[str, float] = {}
            for field_name in numeric_fields:
                try:
                    values[field_name] = _number(
                        row[field_name], f"{location}.{field_name}"
                    )
                except VerificationError:
                    finite_failures += 1
                    values[field_name] = float("nan")
            if not all(math.isfinite(value) for value in values.values()):
                continue

            load_hat = values["predicted_load_kwh"]
            load_actual = values["actual_load_kwh"]
            pv_hat = values["predicted_pv_kwh"]
            pv_actual = values["actual_pv_kwh"]
            grid_value = values["planned_grid_kwh"]
            charge_value = values["planned_charge_kwh"]
            discharge_value = values["planned_discharge_kwh"]
            curtailment = values["predicted_curtailment_kwh"]
            soc_before = values["soc_before_kwh"]
            soc_after = values["soc_after_kwh"]
            emergency = values["actual_emergency_kwh"]
            surplus = values["actual_surplus_kwh"]
            row_price = values["price_yuan_per_kwh"]

            for constrained in (
                load_hat,
                load_actual,
                pv_hat,
                pv_actual,
                grid_value,
                charge_value,
                discharge_value,
                curtailment,
                emergency,
                surplus,
            ):
                if constrained < -abs_tol:
                    negative_failures += 1

            power_limit_excess.add(
                max(0.0, charge_value - energy_limit), location + ".charge"
            )
            power_limit_excess.add(
                max(0.0, discharge_value - energy_limit),
                location + ".discharge",
            )
            if charge_value > abs_tol and discharge_value > abs_tol:
                mutual_exclusion_count += 1
            if emergency > interval_eps and surplus > interval_eps:
                complementarity_count += 1

            soc_lower_violation.add(
                max(0.0, soc_min - soc_before, soc_min - soc_after), location
            )
            soc_upper_violation.add(
                max(0.0, soc_before - soc_max, soc_after - soc_max), location
            )
            expected_soc_after = (
                soc_before + eta_c * charge_value - discharge_value / eta_d
            )
            soc_recursion_error.add(soc_after - expected_soc_after, location)
            if previous_soc_after is None:
                soc_within_day_link_error.add(
                    soc_before - soc_start_json, location + ".start"
                )
            else:
                soc_within_day_link_error.add(
                    soc_before - previous_soc_after, location
                )
            previous_soc_after = soc_after

            predicted_balance_error.add(
                grid_value
                + pv_hat
                + discharge_value
                - load_hat
                - charge_value
                - curtailment,
                location,
            )
            actual_balance_error.add(
                grid_value
                + pv_actual
                + discharge_value
                + emergency
                - load_actual
                - charge_value
                - surplus,
                location,
            )
            plan_cost_error.add(
                values["planned_cost_yuan"] - row_price * grid_value, location
            )
            emergency_cost_error.add(
                values["emergency_cost_yuan"]
                - 5.0 * row_price * emergency,
                location,
            )
            input_price_error.add(
                row_price - price[period_index], location
            )
            input_actual_load_error.add(
                load_actual
                - _number(
                    actual_load_input[source_index][period_index],
                    f"input.actual_load[{source_index}][{period_index}]",
                ),
                location,
            )
            input_actual_pv_error.add(
                pv_actual
                - _number(
                    actual_pv_input[source_index][period_index],
                    f"input.actual_pv[{source_index}][{period_index}]",
                ),
                location,
            )
            json_grid_error.add(
                grid_value
                - _number(grid[period_index], f"{day_text}.grid[{period_index}]"),
                location,
            )
            json_charge_error.add(
                charge_value
                - _number(
                    charge[period_index],
                    f"{day_text}.charge[{period_index}]",
                ),
                location,
            )
            json_discharge_error.add(
                discharge_value
                - _number(
                    discharge[period_index],
                    f"{day_text}.discharge[{period_index}]",
                ),
                location,
            )

            date_grid += grid_value
            date_emergency += emergency
            date_surplus += surplus
            date_plan_cost += values["planned_cost_yuan"]
            date_emergency_cost += values["emergency_cost_yuan"]
            emergency_values.append(emergency)

        if previous_soc_after is not None:
            json_soc_endpoint_error.add(
                previous_soc_after - soc_end_json, f"{day_text}日终"
            )
        json_total_grid_error.add(
            date_grid - _number(day.get("total_grid"), f"{day_text}.total_grid"),
            day_text,
        )
        json_plan_cost_error.add(
            date_plan_cost
            - _number(day.get("plan_cost"), f"{day_text}.plan_cost"),
            day_text,
        )

        expected_intervals = _merge_emergency(
            emergency_values, time_labels, interval_eps
        )
        reconstructed_intervals[day_text] = expected_intervals
        json_intervals = day.get("emergency_intervals")
        _must(
            isinstance(json_intervals, list),
            f"{day_text}.emergency_intervals必须是数组",
        )
        if len(json_intervals) != len(expected_intervals):
            interval_label_mismatches += max(
                1, abs(len(json_intervals) - len(expected_intervals))
            )
        for interval_index, (actual_interval, expected_interval) in enumerate(
            zip(json_intervals, expected_intervals)
        ):
            if str(actual_interval.get("period")) != expected_interval["period"]:
                interval_label_mismatches += 1
            interval_amount_error.add(
                _number(
                    actual_interval.get("amount"),
                    f"{day_text}.emergency_intervals[{interval_index}].amount",
                )
                - float(expected_interval["amount"]),
                f"{day_text}区间{interval_index + 1}",
            )

        daily = daily_by_date[day_text]
        daily_grid_error.add(
            _number(
                daily["planned_grid_kwh"], f"daily {day_text}.planned_grid"
            )
            - date_grid,
            day_text,
        )
        daily_emergency_error.add(
            _number(
                daily["actual_emergency_kwh"],
                f"daily {day_text}.emergency",
            )
            - date_emergency,
            day_text,
        )
        daily_surplus_error.add(
            _number(
                daily["actual_surplus_kwh"], f"daily {day_text}.surplus"
            )
            - date_surplus,
            day_text,
        )
        daily_plan_cost_error.add(
            _number(
                daily["planned_cost_yuan"], f"daily {day_text}.plan_cost"
            )
            - date_plan_cost,
            day_text,
        )
        daily_emergency_cost_error.add(
            _number(
                daily["emergency_cost_yuan"],
                f"daily {day_text}.emergency_cost",
            )
            - date_emergency_cost,
            day_text,
        )
        daily_total_cost_error.add(
            _number(
                daily["actual_total_cost_yuan"],
                f"daily {day_text}.total_cost",
            )
            - date_plan_cost
            - date_emergency_cost,
            day_text,
        )
        terminal_reserve_violation.add(
            max(
                0.0,
                soc_min
                + _number(
                    daily["reserve_kwh"], f"daily {day_text}.reserve"
                )
                - soc_end_json,
            ),
            day_text,
        )
        json_soc_endpoint_error.add(
            _number(daily["soc_start_kwh"], f"daily {day_text}.soc_start")
            - soc_start_json,
            f"daily {day_text}日初",
        )
        json_soc_endpoint_error.add(
            _number(daily["soc_end_kwh"], f"daily {day_text}.soc_end")
            - soc_end_json,
            f"daily {day_text}日终",
        )
        if check_lexicographic_metrics:
            throughput_error.add(
                _number(
                    daily["storage_throughput_kwh"],
                    f"daily {day_text}.storage_throughput",
                )
                - sum(float(value) for value in charge)
                - sum(float(value) for value in discharge),
                day_text,
            )
            peak_grid_error.add(
                _number(
                    daily["peak_grid_kwh"], f"daily {day_text}.peak_grid"
                )
                - max(float(value) for value in grid),
                day_text,
            )
            if str(daily["objective_mode"]) != str(
                diagnostics.get("objective_mode")
            ):
                objective_mode_mismatches += 1

        aggregate_by_date[day_text] = {
            "grid": date_grid,
            "emergency": date_emergency,
            "surplus": date_surplus,
            "plan_cost": date_plan_cost,
            "emergency_cost": date_emergency_cost,
        }
        total_grid += date_grid
        total_emergency += date_emergency
        total_plan_cost += date_plan_cost
        total_emergency_cost += date_emergency_cost

    audit.check(
        "detail_period_sequence",
        period_mismatches == 0,
        f"错误行数={period_mismatches}",
    )
    audit.check(
        "detail_time_mapping",
        time_mismatches == 0,
        f"标题不一致行数={time_mismatches}",
    )
    audit.check(
        "all_numeric_values_finite",
        finite_failures == 0,
        f"非有限或非数值项数={finite_failures}",
    )
    audit.check(
        "nonnegative_energy_variables",
        negative_failures == 0,
        f"低于-{abs_tol:g}的项数={negative_failures}",
    )
    audit.check(
        "charge_discharge_power_limit",
        power_limit_excess.value <= abs_tol,
        f"最大超限={power_limit_excess.value:.12g} kWh，"
        f"位置={power_limit_excess.location}",
    )
    audit.check(
        "charge_discharge_mutual_exclusion",
        mutual_exclusion_count == 0,
        f"同时充放电时段数={mutual_exclusion_count}",
    )
    audit.check(
        "emergency_surplus_complementarity",
        complementarity_count == 0,
        f"同时紧急购电和富余时段数={complementarity_count}",
    )
    audit.check(
        "soc_bounds",
        max(soc_lower_violation.value, soc_upper_violation.value) <= abs_tol,
        f"下界最大违反={soc_lower_violation.value:.12g}，"
        f"上界最大违反={soc_upper_violation.value:.12g} kWh",
    )
    audit.check(
        "soc_recursion",
        soc_recursion_error.value <= abs_tol,
        f"最大误差={soc_recursion_error.value:.12g} kWh，"
        f"位置={soc_recursion_error.location}",
    )
    audit.check(
        "soc_within_day_link",
        soc_within_day_link_error.value <= abs_tol,
        f"最大误差={soc_within_day_link_error.value:.12g} kWh，"
        f"位置={soc_within_day_link_error.location}",
    )
    audit.check(
        "soc_cross_day_link",
        soc_cross_day_error.value <= abs_tol,
        f"最大误差={soc_cross_day_error.value:.12g} kWh，"
        f"位置={soc_cross_day_error.location}",
    )
    audit.check(
        "predicted_energy_balance",
        predicted_balance_error.value <= abs_tol,
        f"最大误差={predicted_balance_error.value:.12g} kWh，"
        f"位置={predicted_balance_error.location}",
    )
    audit.check(
        "actual_energy_balance",
        actual_balance_error.value <= abs_tol,
        f"最大误差={actual_balance_error.value:.12g} kWh，"
        f"位置={actual_balance_error.location}",
    )
    audit.check(
        "period_plan_cost",
        plan_cost_error.value <= abs_tol,
        f"最大误差={plan_cost_error.value:.12g} 元，位置={plan_cost_error.location}",
    )
    audit.check(
        "period_emergency_cost",
        emergency_cost_error.value <= abs_tol,
        f"最大误差={emergency_cost_error.value:.12g} 元，"
        f"位置={emergency_cost_error.location}",
    )
    audit.check(
        "detail_matches_input",
        max(
            input_price_error.value,
            input_actual_load_error.value,
            input_actual_pv_error.value,
        )
        <= abs_tol,
        f"电价/实际负载/实际PV最大误差="
        f"{input_price_error.value:.12g}/{input_actual_load_error.value:.12g}/"
        f"{input_actual_pv_error.value:.12g}",
    )
    audit.check(
        "detail_matches_solution_arrays",
        max(
            json_grid_error.value,
            json_charge_error.value,
            json_discharge_error.value,
        )
        <= abs_tol,
        f"购电/充电/放电最大误差="
        f"{json_grid_error.value:.12g}/{json_charge_error.value:.12g}/"
        f"{json_discharge_error.value:.12g} kWh",
    )
    audit.check(
        "soc_endpoints_match_json_and_daily",
        json_soc_endpoint_error.value <= abs_tol,
        f"最大误差={json_soc_endpoint_error.value:.12g} kWh，"
        f"位置={json_soc_endpoint_error.location}",
    )
    audit.check(
        "daily_totals_match_json",
        max(json_total_grid_error.value, json_plan_cost_error.value) <= abs_tol,
        f"购电量/购电费最大误差={json_total_grid_error.value:.12g} kWh/"
        f"{json_plan_cost_error.value:.12g} 元",
    )
    audit.check(
        "emergency_intervals_merged_continuously",
        interval_label_mismatches == 0
        and interval_amount_error.value <= abs_tol,
        f"标签或数量不一致数={interval_label_mismatches}，"
        f"最大电量误差={interval_amount_error.value:.12g} kWh",
    )
    audit.check(
        "daily_csv_aggregates",
        max(
            daily_grid_error.value,
            daily_emergency_error.value,
            daily_surplus_error.value,
            daily_plan_cost_error.value,
            daily_emergency_cost_error.value,
            daily_total_cost_error.value,
        )
        <= abs_tol,
        "购电/紧急/富余/计划费/紧急费/总费最大误差="
        f"{daily_grid_error.value:.12g}/{daily_emergency_error.value:.12g}/"
        f"{daily_surplus_error.value:.12g}/{daily_plan_cost_error.value:.12g}/"
        f"{daily_emergency_cost_error.value:.12g}/"
        f"{daily_total_cost_error.value:.12g}",
    )
    audit.check(
        "daily_solver_status",
        solver_status_failures == 0,
        f"非最优状态日期数={solver_status_failures}",
    )
    audit.check(
        "terminal_reserve_constraint",
        terminal_reserve_violation.value <= abs_tol,
        f"最大违反={terminal_reserve_violation.value:.12g} kWh，"
        f"日期={terminal_reserve_violation.location}",
    )
    audit.check(
        "lexicographic_metrics",
        not check_lexicographic_metrics
        or (
            max(throughput_error.value, peak_grid_error.value) <= abs_tol
            and objective_mode_mismatches == 0
        ),
        "旧版结果未提供三级目标指标，跳过"
        if not check_lexicographic_metrics
        else f"吞吐量/峰值最大误差={throughput_error.value:.12g}/"
        f"{peak_grid_error.value:.12g}，目标模式错误日期数="
        f"{objective_mode_mismatches}",
    )

    diagnostic_errors = Maximum()
    diagnostic_errors.add(
        _number(
            diagnostics.get("formal_total_grid_kwh"),
            "diagnostics.formal_total_grid_kwh",
        )
        - total_grid,
        "formal_total_grid_kwh",
    )
    diagnostic_errors.add(
        _number(
            diagnostics.get("formal_total_emergency_kwh"),
            "diagnostics.formal_total_emergency_kwh",
        )
        - total_emergency,
        "formal_total_emergency_kwh",
    )
    diagnostic_errors.add(
        _number(
            diagnostics.get("formal_plan_cost_yuan"),
            "diagnostics.formal_plan_cost_yuan",
        )
        - total_plan_cost,
        "formal_plan_cost_yuan",
    )
    diagnostic_errors.add(
        _number(
            diagnostics.get("formal_emergency_cost_yuan"),
            "diagnostics.formal_emergency_cost_yuan",
        )
        - total_emergency_cost,
        "formal_emergency_cost_yuan",
    )
    diagnostic_errors.add(
        _number(
            diagnostics.get("formal_total_cost_yuan"),
            "diagnostics.formal_total_cost_yuan",
        )
        - total_plan_cost
        - total_emergency_cost,
        "formal_total_cost_yuan",
    )
    audit.check(
        "diagnostic_annual_totals",
        diagnostic_errors.value <= abs_tol,
        f"最大误差={diagnostic_errors.value:.12g}，字段={diagnostic_errors.location}",
    )

    _must(workbook_path.is_file(), f"工作簿不存在: {workbook_path}")
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        audit.check(
            "worksheet_names_and_order",
            workbook.sheetnames == REQUIRED_SHEETS,
            f"实际={workbook.sheetnames}",
        )
        _must(
            all(name in workbook.sheetnames for name in REQUIRED_SHEETS),
            "工作簿缺少必需工作表",
        )
        plan = workbook["计划购电量"]
        battery = workbook["充放电量"]
        emergency_sheet = workbook["紧急购电量"]
        for sheet in (plan, battery, emergency_sheet):
            sheet.calculate_dimension(force=True)

        expected_plan_rows = len(days) + 1
        expected_battery_rows = len(days) * 6 + 1
        expected_emergency_rows = (
            sum(len(value) for value in reconstructed_intervals.values()) + 1
        )
        if expected_emergency_rows == 1:
            expected_emergency_rows = 2
        audit.check(
            "worksheet_dimensions",
            plan.max_row == expected_plan_rows
            and plan.max_column == 147
            and battery.max_row == expected_battery_rows
            and battery.max_column == 6
            and emergency_sheet.max_row == expected_emergency_rows
            and emergency_sheet.max_column == 3,
            f"计划={plan.max_row}×{plan.max_column}，"
            f"充放电={battery.max_row}×{battery.max_column}，"
            f"紧急={emergency_sheet.max_row}×{emergency_sheet.max_column}",
        )

        plan_header = list(
            next(
                plan.iter_rows(
                    min_row=1,
                    max_row=1,
                    min_col=1,
                    max_col=147,
                    values_only=True,
                )
            )
        )
        workbook_time_labels = [str(value) for value in plan_header[1:145]]
        workbook_time_ok, workbook_time_detail = _validate_time_labels(
            workbook_time_labels
        )
        audit.check(
            "excel_time_headers",
            workbook_time_ok and workbook_time_labels == time_labels,
            f"{workbook_time_detail}；与input一致="
            f"{workbook_time_labels == time_labels}",
        )

        excel_plan_grid_error = Maximum()
        excel_plan_total_error = Maximum()
        excel_plan_cost_error = Maximum()
        excel_plan_date_mismatches = 0
        plan_rows = plan.iter_rows(
            min_row=2,
            max_row=expected_plan_rows,
            min_col=1,
            max_col=147,
            values_only=True,
        )
        for row_number, (values, day_text) in enumerate(
            zip(plan_rows, formal_expected, strict=True), start=2
        ):
            if _iso_date(values[0], f"计划购电量!A{row_number}") != day_text:
                excel_plan_date_mismatches += 1
            rows = detail_by_date[day_text]
            for period_index in range(N_TIME):
                excel_plan_grid_error.add(
                    _number(
                        values[period_index + 1],
                        f"计划购电量!row{row_number},period{period_index + 1}",
                    )
                    - _number(
                        rows[period_index]["planned_grid_kwh"],
                        f"detail {day_text}#{period_index + 1}.grid",
                    ),
                    f"计划购电量!row{row_number},period{period_index + 1}",
                )
            excel_plan_total_error.add(
                _number(values[145], f"计划购电量!EP{row_number}")
                - aggregate_by_date[day_text]["grid"],
                f"计划购电量!EP{row_number}",
            )
            excel_plan_cost_error.add(
                _number(values[146], f"计划购电量!EQ{row_number}")
                - aggregate_by_date[day_text]["plan_cost"],
                f"计划购电量!EQ{row_number}",
            )
        audit.check(
            "excel_plan_mapping",
            excel_plan_date_mismatches == 0
            and max(
                excel_plan_grid_error.value,
                excel_plan_total_error.value,
                excel_plan_cost_error.value,
            )
            <= abs_tol,
            f"日期错误={excel_plan_date_mismatches}，逐时/合计/费用最大误差="
            f"{excel_plan_grid_error.value:.12g}/{excel_plan_total_error.value:.12g}/"
            f"{excel_plan_cost_error.value:.12g}",
        )

        excel_block_error = Maximum()
        excel_battery_date_mismatches = 0
        excel_battery_label_mismatches = 0
        excel_soc_error = Maximum()
        battery_rows = battery.iter_rows(
            min_row=2,
            max_row=expected_battery_rows,
            min_col=1,
            max_col=6,
            values_only=True,
        )
        for day, day_text in zip(days, formal_expected, strict=True):
            detail_day = detail_by_date[day_text]
            for block in range(6):
                values = next(battery_rows)
                expected_date = day_text if block == 0 else None
                actual_date = (
                    None
                    if values[0] in (None, "")
                    else _iso_date(
                        values[0], f"充放电量 {day_text} block{block + 1}"
                    )
                )
                if actual_date != expected_date:
                    excel_battery_date_mismatches += 1
                if str(values[1]) != BLOCK_LABELS[block]:
                    excel_battery_label_mismatches += 1
                expected_charge = sum(
                    _number(
                        row["planned_charge_kwh"], f"detail {day_text}.charge"
                    )
                    for row in detail_day[
                        block * 24 : (block + 1) * 24
                    ]
                )
                expected_discharge = sum(
                    _number(
                        row["planned_discharge_kwh"],
                        f"detail {day_text}.discharge",
                    )
                    for row in detail_day[
                        block * 24 : (block + 1) * 24
                    ]
                )
                excel_block_error.add(
                    _number(
                        values[2],
                        f"充放电量 {day_text} block{block + 1}.charge",
                    )
                    - expected_charge,
                    f"{day_text} block{block + 1}.charge",
                )
                excel_block_error.add(
                    _number(
                        values[3],
                        f"充放电量 {day_text} block{block + 1}.discharge",
                    )
                    - expected_discharge,
                    f"{day_text} block{block + 1}.discharge",
                )
                expected_time = (
                    "0:00" if block == 0 else "24:00" if block == 1 else None
                )
                expected_soc = (
                    _number(day.get("soc_start"), f"{day_text}.soc_start")
                    if block == 0
                    else _number(day.get("soc_end"), f"{day_text}.soc_end")
                    if block == 1
                    else None
                )
                actual_time = (
                    None if values[4] in (None, "") else str(values[4])
                )
                if actual_time != expected_time:
                    excel_battery_label_mismatches += 1
                if expected_soc is None:
                    if values[5] not in (None, ""):
                        excel_battery_label_mismatches += 1
                else:
                    excel_soc_error.add(
                        _number(
                            values[5],
                            f"充放电量 {day_text} block{block + 1}.soc",
                        )
                        - expected_soc,
                        f"{day_text} block{block + 1}.soc",
                    )
        audit.check(
            "excel_battery_4hour_mapping",
            excel_battery_date_mismatches == 0
            and excel_battery_label_mismatches == 0
            and excel_block_error.value <= abs_tol,
            f"日期错误={excel_battery_date_mismatches}，"
            f"标签错误={excel_battery_label_mismatches}，"
            f"最大汇总误差={excel_block_error.value:.12g} kWh",
        )
        audit.check(
            "excel_soc_start_end_mapping",
            excel_soc_error.value <= abs_tol,
            f"最大误差={excel_soc_error.value:.12g} kWh，"
            f"位置={excel_soc_error.location}",
        )

        expected_emergency: list[tuple[str | None, str, float]] = []
        for day_text in formal_expected:
            intervals = reconstructed_intervals[day_text]
            for interval_index, interval in enumerate(intervals):
                expected_emergency.append(
                    (
                        day_text if interval_index == 0 else None,
                        str(interval["period"]),
                        float(interval["amount"]),
                    )
                )
        if not expected_emergency:
            expected_emergency = [(None, "全年未发生紧急购电", 0.0)]
        excel_emergency_date_mismatches = 0
        excel_emergency_label_mismatches = 0
        excel_emergency_amount_error = Maximum()
        emergency_rows = emergency_sheet.iter_rows(
            min_row=2,
            max_row=expected_emergency_rows,
            min_col=1,
            max_col=3,
            values_only=True,
        )
        for row_number, (values, expected) in enumerate(
            zip(emergency_rows, expected_emergency, strict=True), start=2
        ):
            actual_date = (
                None
                if values[0] in (None, "")
                else _iso_date(values[0], f"紧急购电量!A{row_number}")
            )
            if actual_date != expected[0]:
                excel_emergency_date_mismatches += 1
            if str(values[1]) != expected[1]:
                excel_emergency_label_mismatches += 1
            excel_emergency_amount_error.add(
                _number(values[2], f"紧急购电量!C{row_number}")
                - expected[2],
                f"紧急购电量!C{row_number}",
            )
        audit.check(
            "excel_emergency_interval_mapping",
            excel_emergency_date_mismatches == 0
            and excel_emergency_label_mismatches == 0
            and excel_emergency_amount_error.value <= abs_tol,
            f"日期错误={excel_emergency_date_mismatches}，"
            f"区间错误={excel_emergency_label_mismatches}，"
            f"最大电量误差={excel_emergency_amount_error.value:.12g} kWh",
        )

        excel_error_cells: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if cell.data_type == "e" or (
                        isinstance(value, str) and value.strip() in ERROR_TEXT
                    ):
                        if len(excel_error_cells) < 20:
                            excel_error_cells.append(
                                f"{sheet.title}!{cell.coordinate}={value}"
                            )
                    if isinstance(value, float) and not math.isfinite(value):
                        if len(excel_error_cells) < 20:
                            excel_error_cells.append(
                                f"{sheet.title}!{cell.coordinate}={value}"
                            )
        audit.check(
            "excel_no_error_or_nonfinite_cells",
            not excel_error_cells,
            f"错误单元格={excel_error_cells}",
        )

        plan_rows_count = plan.max_row
        battery_rows_count = battery.max_row
        emergency_rows_count = emergency_sheet.max_row
    finally:
        workbook.close()

    audit.metrics.update(
        {
            "formal_days": len(days),
            "detail_rows": len(detail_rows),
            "daily_rows": len(daily_rows),
            "plan_rows": plan_rows_count,
            "battery_rows": battery_rows_count,
            "emergency_rows": emergency_rows_count,
            "energy_limit_per_period_kwh": energy_limit,
            "formal_total_grid_kwh_recomputed": total_grid,
            "formal_total_emergency_kwh_recomputed": total_emergency,
            "formal_plan_cost_yuan_recomputed": total_plan_cost,
            "formal_emergency_cost_yuan_recomputed": total_emergency_cost,
            "formal_total_cost_yuan_recomputed": total_plan_cost
            + total_emergency_cost,
            "max_predicted_balance_error_kwh": predicted_balance_error.value,
            "max_actual_balance_error_kwh": actual_balance_error.value,
            "max_soc_recursion_error_kwh": soc_recursion_error.value,
            "max_soc_cross_day_error_kwh": soc_cross_day_error.value,
            "max_all_year_soc_cross_day_error_kwh": daily_all_soc_link_error.value,
            "max_terminal_reserve_violation_kwh": terminal_reserve_violation.value,
            "max_excel_grid_error_kwh": excel_plan_grid_error.value,
            "max_excel_block_error_kwh": excel_block_error.value,
            "max_excel_emergency_error_kwh": excel_emergency_amount_error.value,
            "simultaneous_charge_discharge_periods": mutual_exclusion_count,
            "simultaneous_emergency_surplus_periods": complementarity_count,
        }
    )
    return {
        "status": "PASS" if not audit.failures else "FAIL",
        "workbook": str(workbook_path.resolve()),
        "solution": str(solution_path.resolve()),
        "input": str(input_path.resolve()),
        "detail": str(detail_path.resolve()),
        "daily": str(daily_path.resolve()),
        **audit.metrics,
        "checks": audit.checks,
        "failures": audit.failures,
    }


def _write_report(path: Path, report: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="独立核验C题第二问求解、物理约束和result2.xlsx映射"
    )
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--solution", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument(
        "--input",
        type=Path,
        default=None,
        help="默认读取solution同目录的input_data.json",
    )
    parser.add_argument(
        "--detail",
        type=Path,
        default=None,
        help="默认读取solution同目录的question_two_detail.csv",
    )
    parser.add_argument(
        "--daily",
        type=Path,
        default=None,
        help="默认读取solution同目录的question_two_daily.csv",
    )
    parser.add_argument("--abs-tol", type=float, default=1e-6)
    parser.add_argument("--interval-eps", type=float, default=1e-7)
    args = parser.parse_args()

    solution_parent = args.solution.parent
    input_path = _resolve_companion(
        args.input, solution_parent, "input_data.json"
    )
    detail_path = _resolve_companion(
        args.detail, solution_parent, "question_two_detail.csv"
    )
    daily_path = _resolve_companion(
        args.daily, solution_parent, "question_two_daily.csv"
    )

    try:
        report = verify(
            workbook_path=args.workbook,
            solution_path=args.solution,
            input_path=input_path,
            detail_path=detail_path,
            daily_path=daily_path,
            abs_tol=args.abs_tol,
            interval_eps=args.interval_eps,
        )
    except Exception as exc:
        report = {
            "status": "FAIL",
            "workbook": str(args.workbook.resolve()),
            "solution": str(args.solution.resolve()),
            "input": str(input_path.resolve()),
            "detail": str(detail_path.resolve()),
            "daily": str(daily_path.resolve()),
            "fatal_error": f"{type(exc).__name__}: {exc}",
        }
        _write_report(args.report, report)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        raise SystemExit(1) from exc

    _write_report(args.report, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
