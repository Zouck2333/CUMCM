from __future__ import annotations

import argparse
import csv
import json
import math
import time
from dataclasses import dataclass
from datetime import date, datetime, time as clock_time
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import csr_matrix, lil_matrix, vstack


N_TIME = 144
TAUS = (0, 36, 72, 108)
DELTA_H = 1.0 / 6.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_INITIAL = 6000.0
POWER_LIMIT_KW = 5000.0
ENERGY_LIMIT = POWER_LIMIT_KW * DELTA_H
ETA_C = 0.9
ETA_D = 0.9
RISK_ALPHA = 0.90
RISK_WEIGHT = 0.10
PREFIX_WEIGHT = 0.10
SCENARIO_MAX = 60
FORMAL_START = date(2025, 2, 1)
EPS = 1e-8


@dataclass
class InputData:
    dates: list[date]
    price: np.ndarray
    prior_load: np.ndarray
    prior_pv: np.ndarray
    actual_load: np.ndarray
    actual_pv: np.ndarray
    forecast_pv: np.ndarray


@dataclass
class StageSolution:
    grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    curtailment: np.ndarray
    soc: np.ndarray
    emergency: np.ndarray
    objective: float
    throughput: float
    peak_power_kw: float
    status: int
    message: str


def physical_interval_labels() -> list[str]:
    labels: list[str] = []
    for period in range(N_TIME):
        start_minutes = period * 10
        end_minutes = (period + 1) * 10
        start = f"{start_minutes // 60}:{start_minutes % 60:02d}"
        end = (
            "0:00+1"
            if end_minutes == 24 * 60
            else f"{end_minutes // 60}:{end_minutes % 60:02d}"
        )
        labels.append(f"{start}-{end}")
    return labels


def is_workday(value: date) -> bool:
    return value.weekday() < 5


def parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return datetime.fromisoformat(text).date()


def cell_float(value: Any) -> float:
    if value is None:
        raise ValueError("missing numeric cell")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"non-finite numeric cell: {value!r}")
    return result


def load_inputs(attachment1: Path, attachment2: Path, attachment3: Path) -> InputData:
    wb1 = load_workbook(attachment1, data_only=True, read_only=False)
    ws1 = wb1.active
    if ws1.max_row < 145 or ws1.max_column < 4:
        raise ValueError("attachment1 must contain 144 time rows and four columns")
    price = np.asarray(
        [cell_float(ws1.cell(row, 2).value) for row in range(2, 146)],
        dtype=float,
    )
    prior_load = np.asarray(
        [cell_float(ws1.cell(row, 3).value) * DELTA_H for row in range(2, 146)],
        dtype=float,
    )
    prior_pv = np.asarray(
        [
            max(0.0, cell_float(ws1.cell(row, 4).value) * DELTA_H)
            for row in range(2, 146)
        ],
        dtype=float,
    )
    wb1.close()

    wb2 = load_workbook(attachment2, data_only=True, read_only=False)
    if "小区负载" not in wb2.sheetnames or "光伏发电实际功率" not in wb2.sheetnames:
        raise ValueError("attachment2 must contain 小区负载 and 光伏发电实际功率 sheets")
    load_ws = wb2["小区负载"]
    pv_ws = wb2["光伏发电实际功率"]
    if load_ws.max_row < 366 or pv_ws.max_row < 366:
        raise ValueError("attachment2 must contain 365 daily rows")
    if load_ws.max_column < 145 or pv_ws.max_column < 145:
        raise ValueError("attachment2 must contain 144 periods")

    dates: list[date] = []
    actual_load: list[list[float]] = []
    actual_pv: list[list[float]] = []
    for row in range(2, 367):
        load_date = parse_date(load_ws.cell(row, 1).value)
        pv_date = parse_date(pv_ws.cell(row, 1).value)
        if load_date != pv_date:
            raise ValueError(f"attachment2 date mismatch at row {row}")
        dates.append(load_date)
        actual_load.append(
            [
                max(0.0, cell_float(load_ws.cell(row, col).value) * DELTA_H)
                for col in range(2, 146)
            ]
        )
        actual_pv.append(
            [
                max(0.0, cell_float(pv_ws.cell(row, col).value) * DELTA_H)
                for col in range(2, 146)
            ]
        )
    wb2.close()

    if dates[0] != date(2025, 1, 1) or dates[-1] != date(2025, 12, 31):
        raise ValueError("attachment2 dates must cover 2025-01-01 through 2025-12-31")
    if any((dates[i + 1] - dates[i]).days != 1 for i in range(364)):
        raise ValueError("attachment2 dates must be continuous")

    wb3 = load_workbook(attachment3, data_only=True, read_only=False)
    ws3 = wb3.active
    if ws3.max_row < 1461 or ws3.max_column < 26:
        raise ValueError("attachment3 must contain 1460 forecast rows and 24 forecast columns")
    forecast_pv = np.zeros((365, 4, 24), dtype=float)
    current_date: date | None = None
    day_index = -1
    for row in range(2, 1462):
        raw_date = ws3.cell(row, 1).value
        if raw_date not in (None, ""):
            current_date = parse_date(raw_date)
            day_index += 1
            if day_index >= 365:
                raise ValueError("attachment3 contains too many dates")
        if current_date is None or day_index < 0:
            raise ValueError(f"attachment3 row {row} has no date context")
        update_text = str(ws3.cell(row, 2).value).strip()
        if update_text not in {"0:00", "6:00", "12:00", "18:00"}:
            raise ValueError(f"attachment3 row {row} has invalid update time {update_text!r}")
        update_index = {"0:00": 0, "6:00": 1, "12:00": 2, "18:00": 3}[update_text]
        for hour in range(24):
            forecast_pv[day_index, update_index, hour] = max(
                0.0,
                cell_float(ws3.cell(row, 3 + hour).value) * DELTA_H,
            )
    wb3.close()
    if day_index != 364:
        raise ValueError(f"attachment3 expected 365 dates, got {day_index + 1}")

    return InputData(
        dates=dates,
        price=price,
        prior_load=prior_load,
        prior_pv=prior_pv,
        actual_load=np.asarray(actual_load, dtype=float),
        actual_pv=np.asarray(actual_pv, dtype=float),
        forecast_pv=forecast_pv,
    )


def pv_hat_for_day(data: InputData, day: int, k: int) -> np.ndarray:
    tau = TAUS[k]
    result = np.zeros(N_TIME, dtype=float)
    for period in range(tau, N_TIME):
        forecast_hour = (period - tau + 6) // 6
        if forecast_hour < 1 or forecast_hour > 24:
            raise ValueError("forecast horizon index out of range")
        result[period] = data.forecast_pv[day, k, forecast_hour - 1]
    return result


def build_baseline_loader(data: InputData):
    cache: dict[int, np.ndarray] = {}

    def baseline(day: int) -> np.ndarray:
        if day in cache:
            return cache[day]
        if day == 0:
            value = data.prior_load.copy()
            cache[day] = value
            return value
        same_type = [
            index
            for index in range(day)
            if is_workday(data.dates[index]) == is_workday(data.dates[day])
        ]
        selected = sorted(same_type, reverse=True)[:4]
        if len(selected) < 4:
            for index in sorted(range(day), reverse=True):
                if index not in selected:
                    selected.append(index)
                if len(selected) == 4:
                    break
        weights = np.asarray([0.9 ** (day - index) for index in selected], dtype=float)
        weights /= float(np.sum(weights))
        value = np.sum(data.actual_load[selected] * weights[:, None], axis=0)
        cache[day] = value
        return value

    return baseline


def build_load_hat_loader(data: InputData, baseline_loader):
    cache: dict[tuple[int, int], np.ndarray] = {}

    def load_hat(day: int, k: int) -> np.ndarray:
        key = (day, k)
        if key in cache:
            return cache[key]
        base = baseline_loader(day)
        bias = 0.0
        if k > 0:
            start = TAUS[k - 1]
            end = TAUS[k]
            bias = float(
                np.mean(data.actual_load[day, start:end] - base[start:end])
            )
        value = np.maximum(0.0, base + bias)
        cache[key] = value
        return value

    return load_hat


def scenario_bundle(
    data: InputData,
    day: int,
    k: int,
    load_hat_loader,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    tau = TAUS[k]
    horizon = N_TIME - tau
    current_load = load_hat_loader(day, k)
    current_pv = pv_hat_for_day(data, day, k)

    same_type = [
        index
        for index in range(day)
        if is_workday(data.dates[index]) == is_workday(data.dates[day])
    ]
    candidates = sorted(same_type, reverse=True)[:SCENARIO_MAX]
    if len(candidates) < 3:
        for index in sorted(range(day), reverse=True):
            if index not in candidates:
                candidates.append(index)
            if len(candidates) >= min(3, day):
                break

    if not candidates:
        load_scenarios = current_load[tau:][None, :]
        pv_scenarios = current_pv[tau:][None, :]
        probabilities = np.asarray([1.0], dtype=float)
        return current_load, current_pv, load_scenarios, pv_scenarios, probabilities

    residual_load: dict[int, np.ndarray] = {}
    residual_pv: dict[int, np.ndarray] = {}
    for index in candidates:
        historical_load = load_hat_loader(index, k)
        historical_pv = pv_hat_for_day(data, index, k)
        residual_load[index] = data.actual_load[index, tau:] - historical_load[tau:]
        residual_pv[index] = data.actual_pv[index, tau:] - historical_pv[tau:]

    if len(candidates) == 1 or len(candidates) == 2:
        representative = list(candidates)
        base_probabilities = np.full(len(representative), 1.0 / len(representative))
    else:
        ranked = sorted(
            candidates,
            key=lambda index: (float(np.sum(residual_pv[index])), index),
        )
        groups: list[list[int]] = [[], [], []]
        for rank, index in enumerate(ranked):
            group = 1 + int(math.floor(3.0 * rank / len(ranked)))
            groups[group - 1].append(index)
        representative = []
        base_probabilities = []
        for group in groups:
            middle = (len(group) - 1) // 2
            representative.append(group[middle])
            base_probabilities.append(float(len(group)) / float(len(ranked)))
        base_probabilities = np.asarray(base_probabilities, dtype=float)

    if k > 0 and candidates:
        prefix_indices = np.arange(0, tau)
        historical_errors = np.vstack(
            [
                data.actual_load[index, prefix_indices]
                - load_hat_loader(index, k)[prefix_indices]
                for index in candidates
            ]
        )
        current_errors = (
            data.actual_load[day, prefix_indices] - current_load[prefix_indices]
        )
        scale = np.std(historical_errors, axis=0, ddof=0) + 1e-9
        distances = np.mean(
            ((historical_errors - current_errors[None, :]) / scale[None, :]) ** 2,
            axis=1,
        )
        distance_by_day = {
            index: float(distances[position])
            for position, index in enumerate(candidates)
        }
        log_weights = np.log(np.maximum(base_probabilities, 1e-300)) - (
            PREFIX_WEIGHT
            * np.asarray([distance_by_day[index] for index in representative])
        )
        log_weights -= float(np.max(log_weights))
        weights = np.exp(log_weights)
        probabilities = weights / float(np.sum(weights))
    else:
        probabilities = base_probabilities

    load_scenarios = np.vstack(
        [
            np.maximum(0.0, current_load[tau:] + residual_load[index])
            for index in representative
        ]
    )
    pv_scenarios = np.vstack(
        [
            np.maximum(0.0, current_pv[tau:] + residual_pv[index])
            for index in representative
        ]
    )
    return (
        current_load,
        current_pv,
        load_scenarios,
        pv_scenarios,
        np.asarray(probabilities, dtype=float),
    )


def compute_reserve(
    data: InputData,
    day: int,
    k: int,
    load_hat_loader,
    quantile: float,
    eta_discharge: float,
) -> float:
    if quantile <= 0.0:
        return 0.0
    tau = TAUS[k]
    same_type = [
        index
        for index in range(day)
        if is_workday(data.dates[index]) == is_workday(data.dates[day])
    ]
    candidates = sorted(same_type, reverse=True)[:SCENARIO_MAX]
    if len(candidates) < 5:
        return 0.0
    requirements: list[float] = []
    for index in candidates:
        historical_load = load_hat_loader(index, k)
        historical_pv = pv_hat_for_day(data, index, k)
        load_error = data.actual_load[index, tau:] - historical_load[tau:]
        pv_error = data.actual_pv[index, tau:] - historical_pv[tau:]
        requirements.append(
            float(np.sum(np.maximum(0.0, load_error - pv_error)))
        )
    raw = float(np.quantile(np.asarray(requirements), quantile, method="linear"))
    raw /= eta_discharge
    return max(0.0, min(SOC_MAX - SOC_MIN, raw))


def solve_stage(
    price: np.ndarray,
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    probabilities: np.ndarray,
    *,
    k: int,
    soc_start: float,
    initial_grid: np.ndarray | None,
    eta_charge: float,
    eta_discharge: float,
    risk_alpha: float,
    risk_weight: float,
    mip_rel_gap: float = 1e-5,
    reserve: float = 0.0,
    lexicographic: bool = False,
) -> StageSolution:
    horizon = len(price)
    scenario_count = load_scenarios.shape[0]
    grid_start = 0
    charge_start = horizon
    discharge_start = 2 * horizon
    curtail_start = 3 * horizon
    soc_start_idx = 4 * horizon
    emergency_start = soc_start_idx + horizon + 1
    u_start = emergency_start + scenario_count * horizon
    v_start = u_start + horizon
    zeta_idx = v_start + horizon
    a_start = zeta_idx + 1
    mode_start = a_start + scenario_count
    peak_idx = mode_start + horizon
    variable_count = peak_idx + 1

    objective = np.zeros(variable_count, dtype=float)
    for period in range(horizon):
        if k == 0:
            objective[grid_start + period] = price[period]
        else:
            objective[u_start + period] = 1.5 * price[period]
            objective[v_start + period] = 0.5 * price[period]
    for scenario in range(scenario_count):
        objective[
            emergency_start + scenario * horizon : emergency_start
            + (scenario + 1) * horizon
        ] = 5.0 * price * probabilities[scenario]
        objective[a_start + scenario] = (
            risk_weight / (1.0 - risk_alpha) * probabilities[scenario]
        )
    objective[zeta_idx] = risk_weight

    lower = np.zeros(variable_count, dtype=float)
    upper = np.full(variable_count, np.inf, dtype=float)
    upper[charge_start : charge_start + horizon] = ENERGY_LIMIT
    upper[discharge_start : discharge_start + horizon] = ENERGY_LIMIT
    upper[curtail_start : curtail_start + horizon] = pv_hat
    lower[soc_start_idx : soc_start_idx + horizon + 1] = SOC_MIN
    upper[soc_start_idx : soc_start_idx + horizon + 1] = SOC_MAX
    lower[soc_start_idx] = soc_start
    upper[soc_start_idx] = soc_start
    lower[zeta_idx] = -np.inf
    upper[mode_start : mode_start + horizon] = 1.0
    if reserve > 0.0:
        lower[soc_start_idx + horizon] = max(
            lower[soc_start_idx + horizon],
            SOC_MIN + reserve,
        )

    equalities: list[dict[int, float]] = []
    equality_rhs: list[float] = []
    inequalities: list[dict[int, float]] = []
    inequality_rhs: list[float] = []

    for period in range(horizon):
        equalities.append(
            {
                grid_start + period: 1.0,
                discharge_start + period: 1.0,
                charge_start + period: -1.0,
                curtail_start + period: -1.0,
            }
        )
        equality_rhs.append(float(load_hat[period] - pv_hat[period]))

    for period in range(horizon):
        row = {
            soc_start_idx + period + 1: 1.0,
            soc_start_idx + period: -1.0,
            charge_start + period: -eta_charge,
            discharge_start + period: 1.0 / eta_discharge,
        }
        equalities.append(row)
        equality_rhs.append(0.0)

    if k > 0:
        assert initial_grid is not None
        for period in range(horizon):
            equalities.append(
                {
                    grid_start + period: 1.0,
                    u_start + period: -1.0,
                    v_start + period: 1.0,
                }
            )
            equality_rhs.append(float(initial_grid[period]))

    for scenario in range(scenario_count):
        for period in range(horizon):
            inequalities.append(
                {
                    grid_start + period: -1.0,
                    charge_start + period: 1.0,
                    discharge_start + period: -1.0,
                    emergency_start + scenario * horizon + period: -1.0,
                }
            )
            inequality_rhs.append(
                float(
                    pv_scenarios[scenario, period]
                    - load_scenarios[scenario, period]
                )
            )

    for period in range(horizon):
        inequalities.append(
            {
                charge_start + period: 1.0,
                mode_start + period: -ENERGY_LIMIT,
            }
        )
        inequality_rhs.append(0.0)
        inequalities.append(
            {
                discharge_start + period: 1.0,
                mode_start + period: ENERGY_LIMIT,
            }
        )
        inequality_rhs.append(ENERGY_LIMIT)

    for scenario in range(scenario_count):
        row: dict[int, float] = {
            a_start + scenario: -1.0,
            zeta_idx: -1.0,
        }
        if k == 0:
            for period in range(horizon):
                row[grid_start + period] = price[period]
                row[emergency_start + scenario * horizon + period] = (
                    5.0 * price[period]
                )
        else:
            for period in range(horizon):
                row[u_start + period] = 1.5 * price[period]
                row[v_start + period] = 0.5 * price[period]
                row[emergency_start + scenario * horizon + period] = (
                    5.0 * price[period]
                )
        inequalities.append(row)
        inequality_rhs.append(0.0)

    for scenario in range(scenario_count):
        for period in range(horizon):
            inequalities.append(
                {
                    grid_start + period: 1.0,
                    emergency_start + scenario * horizon + period: 1.0,
                    peak_idx: -DELTA_H,
                }
            )
            inequality_rhs.append(0.0)

    a_eq = lil_matrix((len(equalities), variable_count), dtype=float)
    for row_index, row in enumerate(equalities):
        for column, value in row.items():
            a_eq[row_index, column] = value
    a_ub = lil_matrix((len(inequalities), variable_count), dtype=float)
    for row_index, row in enumerate(inequalities):
        for column, value in row.items():
            a_ub[row_index, column] = value

    integrality = np.zeros(variable_count, dtype=np.uint8)
    integrality[mode_start : mode_start + horizon] = 1
    milp_options = {
        "presolve": True,
        "mip_rel_gap": mip_rel_gap,
    }
    result = milp(
        c=objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=[
            LinearConstraint(
                a_ub.tocsr(),
                np.full(len(inequality_rhs), -np.inf, dtype=float),
                np.asarray(inequality_rhs, dtype=float),
            ),
            LinearConstraint(
                a_eq.tocsr(),
                np.asarray(equality_rhs, dtype=float),
                np.asarray(equality_rhs, dtype=float),
            ),
        ],
        options=milp_options,
    )
    if result.x is None or result.status != 0:
        raise RuntimeError(f"stage {k} failed: status={result.status}, {result.message}")

    primary_optimum = float(result.fun)
    if lexicographic:
        cost_tolerance = max(1e-6, abs(primary_optimum) * 1e-7)
        throughput_objective = np.zeros(variable_count, dtype=float)
        throughput_objective[charge_start : charge_start + horizon] = 1.0
        throughput_objective[discharge_start : discharge_start + horizon] = 1.0
        stage_two_ub = vstack(
            [a_ub.tocsr(), csr_matrix(objective.reshape(1, -1))],
            format="csr",
        )
        stage_two_rhs = np.append(
            np.asarray(inequality_rhs, dtype=float),
            primary_optimum + cost_tolerance,
        )
        stage_two = milp(
            c=throughput_objective,
            integrality=integrality,
            bounds=Bounds(lower, upper),
            constraints=[
                LinearConstraint(
                    stage_two_ub,
                    np.full(len(stage_two_rhs), -np.inf, dtype=float),
                    stage_two_rhs,
                ),
                LinearConstraint(
                    a_eq.tocsr(),
                    np.asarray(equality_rhs, dtype=float),
                    np.asarray(equality_rhs, dtype=float),
                ),
            ],
            options=milp_options,
        )
        if stage_two.x is None or stage_two.status != 0:
            raise RuntimeError(
                f"stage {k} lexicographic throughput failed: "
                f"status={stage_two.status}, {stage_two.message}"
            )

        throughput_optimum = float(stage_two.fun)
        throughput_tolerance = max(1e-6, abs(throughput_optimum) * 1e-9)
        stage_three_ub = vstack(
            [
                stage_two_ub,
                csr_matrix(throughput_objective.reshape(1, -1)),
            ],
            format="csr",
        )
        stage_three_rhs = np.append(
            stage_two_rhs,
            throughput_optimum + throughput_tolerance,
        )
        peak_objective = np.zeros(variable_count, dtype=float)
        peak_objective[peak_idx] = 1.0
        stage_three = milp(
            c=peak_objective,
            integrality=integrality,
            bounds=Bounds(lower, upper),
            constraints=[
                LinearConstraint(
                    stage_three_ub,
                    np.full(len(stage_three_rhs), -np.inf, dtype=float),
                    stage_three_rhs,
                ),
                LinearConstraint(
                    a_eq.tocsr(),
                    np.asarray(equality_rhs, dtype=float),
                    np.asarray(equality_rhs, dtype=float),
                ),
            ],
            options=milp_options,
        )
        if stage_three.x is None or stage_three.status != 0:
            raise RuntimeError(
                f"stage {k} lexicographic peak failed: "
                f"status={stage_three.status}, {stage_three.message}"
            )
        result = stage_three

    values = np.asarray(result.x, dtype=float)
    values[np.abs(values) < 1e-8] = 0.0
    grid = np.maximum(0.0, values[grid_start : grid_start + horizon])
    charge = np.maximum(0.0, values[charge_start : charge_start + horizon])
    discharge = np.maximum(0.0, values[discharge_start : discharge_start + horizon])
    soc = values[soc_start_idx : soc_start_idx + horizon + 1]
    emergency = np.maximum(
        0.0,
        values[
            emergency_start : emergency_start + scenario_count * horizon
        ].reshape(scenario_count, horizon),
    )
    curtailment = np.maximum(
        0.0,
        grid + pv_hat + discharge - load_hat - charge,
    )
    throughput = float(np.sum(charge) + np.sum(discharge))
    peak_power_kw = float(max(0.0, values[peak_idx]))
    return StageSolution(
        grid=grid,
        charge=charge,
        discharge=discharge,
        curtailment=curtailment,
        soc=soc,
        emergency=emergency,
        objective=primary_optimum,
        throughput=throughput,
        peak_power_kw=peak_power_kw,
        status=int(result.status),
        message=str(result.message),
    )


def merge_intervals(
    emergency: np.ndarray,
    time_labels: list[str],
    threshold: float = 1e-8,
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    start: int | None = None
    for index, value in enumerate(emergency):
        active = float(value) > threshold
        if active and start is None:
            start = index
        if start is not None and (not active or index == len(emergency) - 1):
            end = index if active and index == len(emergency) - 1 else index - 1
            first = time_labels[start].split("-", 1)[0]
            last = time_labels[end].split("-", 1)[-1]
            result.append(
                {
                    "period": f"{first}-{last}",
                    "amount": float(np.sum(emergency[start : end + 1])),
                }
            )
            start = None
    return result


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_result_workbook(
    template_path: Path,
    output_path: Path,
    days: list[dict[str, Any]],
    time_labels: list[str],
) -> None:
    workbook = load_workbook(template_path)
    names = workbook.sheetnames
    plan_sheet = workbook[names[0]]
    actual_sheet = workbook[names[1]]
    battery_sheet = workbook[names[2]]
    emergency_sheet = workbook[names[3]]

    for sheet in (plan_sheet, actual_sheet):
        for column, label in enumerate(time_labels, start=2):
            sheet.cell(1, column, label)

    for row_index, day in enumerate(days, start=2):
        plan_sheet.cell(row_index, 1, day["date"])
        actual_sheet.cell(row_index, 1, day["date"])
        for period in range(N_TIME):
            plan_sheet.cell(row_index, 2 + period, day["initial_grid"][period])
            actual_sheet.cell(row_index, 2 + period, day["grid"][period])
        plan_sheet.cell(row_index, 146, day["total_initial_grid"])
        plan_sheet.cell(row_index, 147, day["plan_cost"])
        actual_sheet.cell(row_index, 146, day["total_grid"])
        actual_sheet.cell(row_index, 147, day["total_cost"])

    battery_row = 2
    for day in days:
        for block in range(6):
            for column in range(1, 7):
                battery_sheet.cell(battery_row, column, "")
            if block == 0:
                battery_sheet.cell(battery_row, 1, day["date"])
            battery_sheet.cell(battery_row, 2, f"{4 * block}:00-{4 * (block + 1)}:00")
            start = 24 * block
            battery_sheet.cell(
                battery_row,
                3,
                float(np.sum(day["charge"][start : start + 24])),
            )
            battery_sheet.cell(
                battery_row,
                4,
                float(np.sum(day["discharge"][start : start + 24])),
            )
            if block == 0:
                battery_sheet.cell(battery_row, 5, "0:00")
                battery_sheet.cell(battery_row, 6, day["soc_start"])
            elif block == 1:
                battery_sheet.cell(battery_row, 5, "24:00")
                battery_sheet.cell(battery_row, 6, day["soc_end"])
            battery_row += 1
    if battery_sheet.max_row >= battery_row:
        battery_sheet.delete_rows(battery_row, battery_sheet.max_row - battery_row + 1)

    emergency_row = 2
    for day in days:
        intervals = day["emergency_intervals"]
        if not intervals:
            emergency_sheet.cell(emergency_row, 1, day["date"])
            emergency_row += 1
            continue
        for index, interval in enumerate(intervals):
            if index == 0:
                emergency_sheet.cell(emergency_row, 1, day["date"])
            emergency_sheet.cell(emergency_row, 2, interval["period"])
            emergency_sheet.cell(emergency_row, 3, interval["amount"])
            emergency_row += 1
    if emergency_sheet.max_row >= emergency_row:
        emergency_sheet.delete_rows(
            emergency_row,
            emergency_sheet.max_row - emergency_row + 1,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def solve_year(
    data: InputData,
    *,
    eta_charge: float,
    eta_discharge: float,
    risk_alpha: float,
    risk_weight: float,
    max_days: int | None,
    zero_only: bool = False,
    reserve_quantile: float = 0.0,
    lexicographic: bool = False,
    mip_rel_gap: float = 1e-5,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    baseline_loader = build_baseline_loader(data)
    load_hat_loader = build_load_hat_loader(data, baseline_loader)
    time_labels = physical_interval_labels()
    formal_days: list[dict[str, Any]] = []
    daily_rows: list[dict[str, Any]] = []
    total_overlap = 0.0
    max_balance_residual = 0.0
    limit = len(data.dates) if max_days is None else min(max_days, len(data.dates))
    soc_initial = SOC_INITIAL

    for day_index in range(limit):
        current_date = data.dates[day_index]
        initial_grid = np.zeros(N_TIME, dtype=float)
        final_grid = np.zeros(N_TIME, dtype=float)
        final_charge = np.zeros(N_TIME, dtype=float)
        final_discharge = np.zeros(N_TIME, dtype=float)
        soc_path = np.zeros(N_TIME + 1, dtype=float)
        soc_path[0] = soc_initial
        current_soc = soc_initial
        stage_objectives: list[float] = []

        for k in range(1 if zero_only else 4):
            tau = TAUS[k]
            horizon = N_TIME - tau
            current_load, current_pv, load_scenarios, pv_scenarios, probabilities = (
                scenario_bundle(data, day_index, k, load_hat_loader)
            )
            fixed_initial = None if k == 0 else initial_grid[tau:]
            reserve = compute_reserve(
                data,
                day_index,
                k,
                load_hat_loader,
                reserve_quantile,
                eta_discharge,
            )
            stage = solve_stage(
                data.price[tau:],
                current_load[tau:],
                current_pv[tau:],
                load_scenarios,
                pv_scenarios,
                probabilities,
                k=k,
                soc_start=current_soc,
                initial_grid=fixed_initial,
                eta_charge=eta_charge,
                eta_discharge=eta_discharge,
                risk_alpha=risk_alpha,
                risk_weight=risk_weight,
                mip_rel_gap=mip_rel_gap,
                reserve=reserve,
                lexicographic=lexicographic,
            )
            stage_objectives.append(stage.objective)
            if k == 0:
                initial_grid[:] = stage.grid
            if zero_only:
                final_grid[:] = stage.grid
                final_charge[:] = stage.charge
                final_discharge[:] = stage.discharge
                for period in range(N_TIME):
                    soc_path[period + 1] = (
                        soc_path[period]
                        + eta_charge * stage.charge[period]
                        - stage.discharge[period] / eta_discharge
                    )
                current_soc = float(soc_path[-1])
            else:
                execute_end = tau + 36
                final_grid[tau:execute_end] = stage.grid[:36]
                final_charge[tau:execute_end] = stage.charge[:36]
                final_discharge[tau:execute_end] = stage.discharge[:36]
                for period in range(tau, execute_end):
                    local = period - tau
                    soc_path[period + 1] = (
                        soc_path[period]
                        + eta_charge * stage.charge[local]
                        - stage.discharge[local] / eta_discharge
                    )
                current_soc = float(soc_path[execute_end])

        actual_shortfall = (
            data.actual_load[day_index]
            + final_charge
            - final_grid
            - data.actual_pv[day_index]
            - final_discharge
        )
        actual_emergency = np.maximum(0.0, actual_shortfall)
        actual_surplus = np.maximum(0.0, -actual_shortfall)
        plan_cost = float(np.sum(data.price * initial_grid))
        adjustment_cost = float(
            np.sum(
                1.5
                * data.price
                * np.maximum(0.0, final_grid - initial_grid)
                + 0.5
                * data.price
                * np.maximum(0.0, initial_grid - final_grid)
            )
        )
        emergency_cost = float(np.sum(5.0 * data.price * actual_emergency))
        total_cost = plan_cost + adjustment_cost + emergency_cost
        actual_peak_kw = float(np.max((final_grid + actual_emergency) / DELTA_H))
        balance_residual = (
            final_grid
            + data.actual_pv[day_index]
            + final_discharge
            + actual_emergency
            - data.actual_load[day_index]
            - final_charge
            - actual_surplus
        )
        max_balance_residual = max(
            max_balance_residual,
            float(np.max(np.abs(balance_residual))),
        )
        overlap = float(np.max(np.minimum(final_charge, final_discharge)))
        total_overlap = max(total_overlap, overlap)

        if not (
            SOC_MIN - 1e-5
            <= float(soc_path[-1])
            <= SOC_MAX + 1e-5
        ):
            raise RuntimeError(
                f"{current_date.isoformat()} terminal SOC out of bounds: {soc_path[-1]}"
            )
        soc_initial = float(soc_path[-1])

        if current_date >= FORMAL_START:
            formal_days.append(
                {
                    "date": current_date.isoformat(),
                    "initial_grid": initial_grid.tolist(),
                    "grid": final_grid.tolist(),
                    "charge": final_charge.tolist(),
                    "discharge": final_discharge.tolist(),
                    "soc_path": soc_path.tolist(),
                    "soc_start": float(soc_path[0]),
                    "soc_end": float(soc_path[-1]),
                    "total_initial_grid": float(np.sum(initial_grid)),
                    "total_grid": float(np.sum(final_grid)),
                    "plan_cost": plan_cost,
                    "adjustment_cost": adjustment_cost,
                    "emergency_cost": emergency_cost,
                    "total_cost": total_cost,
                    "actual_peak_kw": actual_peak_kw,
                    "actual_emergency": actual_emergency.tolist(),
                    "actual_surplus": actual_surplus.tolist(),
                    "actual_load": data.actual_load[day_index].tolist(),
                    "actual_pv": data.actual_pv[day_index].tolist(),
                    "emergency_intervals": merge_intervals(
                        actual_emergency,
                        time_labels,
                    ),
                }
            )
        daily_rows.append(
            {
                "date": current_date.isoformat(),
                "soc_start": float(soc_path[0]),
                "soc_end": float(soc_path[-1]),
                "initial_grid": float(np.sum(initial_grid)),
                "final_grid": float(np.sum(final_grid)),
                "plan_cost": plan_cost,
                "adjustment_cost": adjustment_cost,
                "emergency_cost": emergency_cost,
                "total_cost": total_cost,
                "actual_peak_kw": actual_peak_kw,
                "emergency_kwh": float(np.sum(actual_emergency)),
                "surplus_kwh": float(np.sum(actual_surplus)),
                "max_charge_discharge_overlap": overlap,
                "stage_objectives": json.dumps(stage_objectives),
            }
        )
        if day_index == 0 or (day_index + 1) % 10 == 0 or day_index + 1 == limit:
            print(
                f"[{day_index + 1:3d}/{limit}] {current_date.isoformat()} "
                f"SOC={soc_path[-1]:.3f} total={total_cost:.2f}",
                flush=True,
            )

    formal_rows = [row for row in daily_rows if row["date"] >= FORMAL_START.isoformat()]
    diagnostics = {
        "model": (
            "zero-hour-only three-branch stochastic MILP with free terminal SOC and cross-day continuity"
            if zero_only
            else "four-update rolling three-branch stochastic MPC with free terminal SOC and cross-day continuity"
        ),
        "strategy": "zero_hour_only" if zero_only else "four_update_rolling",
        "formal_start": FORMAL_START.isoformat(),
        "solver": "scipy.optimize.milp / HiGHS",
        "eta_charge": eta_charge,
        "eta_discharge": eta_discharge,
        "risk_alpha": risk_alpha,
        "risk_weight": risk_weight,
        "prefix_weight": PREFIX_WEIGHT,
        "reserve_quantile": reserve_quantile,
        "lexicographic": lexicographic,
        "mip_rel_gap": mip_rel_gap,
        "initial_soc_2025_01_01": SOC_INITIAL,
        "terminal_soc_mode": "free within capacity bounds; carried to next day",
        "soc_min": SOC_MIN,
        "soc_max": SOC_MAX,
        "days_solved": limit,
        "formal_days": len(formal_rows),
        "max_balance_residual_kwh": max_balance_residual,
        "max_charge_discharge_overlap_kwh": total_overlap,
        "formal_total_initial_grid_kwh": float(
            sum(row["initial_grid"] for row in formal_rows)
        ),
        "formal_total_final_grid_kwh": float(
            sum(row["final_grid"] for row in formal_rows)
        ),
        "formal_plan_cost_yuan": float(sum(row["plan_cost"] for row in formal_rows)),
        "formal_adjustment_cost_yuan": float(
            sum(row["adjustment_cost"] for row in formal_rows)
        ),
        "formal_emergency_cost_yuan": float(
            sum(row["emergency_cost"] for row in formal_rows)
        ),
        "formal_total_cost_yuan": float(sum(row["total_cost"] for row in formal_rows)),
        "formal_emergency_kwh": float(
            sum(row["emergency_kwh"] for row in formal_rows)
        ),
        "formal_max_actual_peak_kw": float(
            max((row["actual_peak_kw"] for row in formal_rows), default=0.0)
        ),
    }
    return formal_days, daily_rows, diagnostics


def main() -> None:
    parser = argparse.ArgumentParser(description="C题第三问滚动随机 MPC 求解器")
    parser.add_argument("--attachment1", type=Path, required=True)
    parser.add_argument("--attachment2", type=Path, required=True)
    parser.add_argument("--attachment3", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--eta-charge", type=float, default=ETA_C)
    parser.add_argument("--eta-discharge", type=float, default=ETA_D)
    parser.add_argument("--risk-alpha", type=float, default=RISK_ALPHA)
    parser.add_argument("--risk-weight", type=float, default=RISK_WEIGHT)
    parser.add_argument("--reserve-quantile", type=float, default=0.0)
    parser.add_argument("--lexicographic", action="store_true")
    parser.add_argument("--mip-rel-gap", type=float, default=1e-5)
    parser.add_argument("--zero-only", action="store_true")
    parser.add_argument("--max-days", type=int, default=None)
    args = parser.parse_args()

    started = time.perf_counter()
    data = load_inputs(args.attachment1, args.attachment2, args.attachment3)
    formal_days, daily_rows, diagnostics = solve_year(
        data,
        eta_charge=args.eta_charge,
        eta_discharge=args.eta_discharge,
        risk_alpha=args.risk_alpha,
        risk_weight=args.risk_weight,
        max_days=args.max_days,
        zero_only=args.zero_only,
        reserve_quantile=args.reserve_quantile,
        lexicographic=args.lexicographic,
        mip_rel_gap=args.mip_rel_gap,
    )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    solution_path = args.output_dir / "question_three_solution.json"
    solution_path.write_text(
        json.dumps(
            {
                "diagnostics": diagnostics,
                "time_labels": physical_interval_labels(),
                "days": formal_days,
            },
            ensure_ascii=False,
            allow_nan=False,
        ),
        encoding="utf-8",
    )
    write_csv(args.output_dir / "question_three_daily.csv", daily_rows)
    detail_rows: list[dict[str, Any]] = []
    for day in formal_days:
        for period in range(N_TIME):
            detail_rows.append(
                {
                    "date": day["date"],
                    "period": period + 1,
                    "time": physical_interval_labels()[period],
                    "price": float(data.price[period]),
                    "initial_grid": day["initial_grid"][period],
                    "grid": day["grid"][period],
                    "charge": day["charge"][period],
                    "discharge": day["discharge"][period],
                    "soc": day["soc_path"][period + 1],
                    "actual_load": day["actual_load"][period],
                    "actual_pv": day["actual_pv"][period],
                    "emergency": day["actual_emergency"][period],
                    "surplus": day["actual_surplus"][period],
                }
            )
    write_csv(args.output_dir / "question_three_detail.csv", detail_rows)
    write_result_workbook(
        args.template,
        args.output_dir / "result3.xlsx",
        formal_days,
        physical_interval_labels(),
    )
    elapsed = time.perf_counter() - started
    print(json.dumps(diagnostics, ensure_ascii=False, indent=2))
    print(f"elapsed_seconds={elapsed:.3f}")


if __name__ == "__main__":
    main()
