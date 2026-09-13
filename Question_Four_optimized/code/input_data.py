"""Read attachments 2, 3, and 4 into the in-memory Question Four model."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
import math
import sys

import numpy as np

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    if exc.name != "openpyxl":
        raise
    bundled = (
        Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime"
        / "dependencies" / "python" / "Lib" / "site-packages"
    )
    if not (bundled / "openpyxl").is_dir():
        raise ModuleNotFoundError("openpyxl is required to read the attachments") from exc
    sys.path.append(str(bundled))
    from openpyxl import load_workbook

from .model import DayData, LAST_DAY, N_TIME


FIRST_DAY = date(2025, 1, 1)
DAY_COUNT = (LAST_DAY - FIRST_DAY).days + 1
RELEASE_MINUTES = (0, 360, 720, 1080)


@dataclass(frozen=True)
class LoadedInputs:
    days: list[DayData]
    template_4_2: Path
    template_4_3: Path


def _number(value: object, label: str, *, positive: bool = False) -> float:
    if value is None or isinstance(value, bool):
        raise ValueError(f"{label}: missing numeric value")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: invalid numeric value {value!r}") from exc
    if not math.isfinite(result) or (result <= 0 if positive else result < 0):
        raise ValueError(f"{label}: expected {'positive' if positive else 'nonnegative'} finite value")
    return result


def _date(value: object, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        parts = value.strip().split("-")
        if len(parts) == 3:
            try:
                return date(*(int(part) for part in parts))
            except ValueError:
                pass
    raise ValueError(f"{label}: invalid date {value!r}")


def _minute(value: object, label: str) -> int:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"{label}: seconds are not permitted")
        return value.hour * 60 + value.minute
    if isinstance(value, str):
        parts = value.strip().split(":")
        if len(parts) == 2:
            try:
                hour, minute = map(int, parts)
            except ValueError:
                pass
            else:
                if 0 <= hour < 24 and 0 <= minute < 60:
                    return hour * 60 + minute
    raise ValueError(f"{label}: invalid time {value!r}")


def _check_right_endpoint_header(sheet, label: str) -> None:
    if sheet.max_column < N_TIME + 1:
        raise ValueError(f"{label}: expected 144 interval columns")
    header = next(sheet.iter_rows(
        min_row=1, max_row=1, min_col=2, max_col=N_TIME + 1, values_only=True
    ))
    for t, value in enumerate(header):
        expected = (t + 1) * 10
        if expected == 1440 and isinstance(value, str) and value.strip() in {"24:00", "0:00+1"}:
            continue
        if _minute(value, f"{label} header {t}") != expected:
            raise ValueError(f"{label}: interval {t} has the wrong right endpoint")


def _read_daily_sheet(sheet, label: str, *, positive: bool = False) -> np.ndarray:
    if sheet.max_row != DAY_COUNT + 1:
        raise ValueError(f"{label}: expected {DAY_COUNT} dates")
    _check_right_endpoint_header(sheet, label)
    values = np.empty((DAY_COUNT, N_TIME))
    rows = sheet.iter_rows(
        min_row=2, max_row=DAY_COUNT + 1,
        min_col=1, max_col=N_TIME + 1, values_only=True,
    )
    for i, row in enumerate(rows):
        expected_day = FIRST_DAY + timedelta(days=i)
        if _date(row[0], f"{label} row {i+2}") != expected_day:
            raise ValueError(f"{label}: dates are not consecutive at row {i+2}")
        for t, value in enumerate(row[1:]):
            values[i, t] = _number(value, f"{label} row {i+2} interval {t}", positive=positive)
    return values


def _read_actuals(path: Path) -> tuple[np.ndarray, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) < 2:
            raise ValueError("attachment 2 must contain load and PV worksheets")
        load_kw = _read_daily_sheet(workbook.worksheets[0], "attachment 2 load")
        pv_kw = _read_daily_sheet(workbook.worksheets[1], "attachment 2 PV")
        return load_kw / 6.0, pv_kw / 6.0
    finally:
        workbook.close()


def _read_prices(path: Path) -> np.ndarray:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return _read_daily_sheet(workbook.active, "attachment 4 price", positive=True)
    finally:
        workbook.close()


def _read_pv_forecasts(path: Path) -> np.ndarray:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_row != 1 + DAY_COUNT * 4 or sheet.max_column < 26:
            raise ValueError("attachment 3 must have four 24-hour releases per day")
        forecasts = np.empty((DAY_COUNT, 4, 24))
        rows = sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=26, values_only=True)
        for i, row in enumerate(rows):
            day_index, stage = divmod(i, 4)
            expected_day = FIRST_DAY + timedelta(days=day_index)
            if stage == 0:
                if _date(row[0], f"attachment 3 row {i+2}") != expected_day:
                    raise ValueError(f"attachment 3: wrong date at row {i+2}")
            elif row[0] not in (None, ""):
                raise ValueError(f"attachment 3: duplicate date at row {i+2}")
            if _minute(row[1], f"attachment 3 row {i+2}") != RELEASE_MINUTES[stage]:
                raise ValueError(f"attachment 3: wrong release time at row {i+2}")
            for j, value in enumerate(row[2:]):
                forecasts[day_index, stage, j] = _number(
                    value, f"attachment 3 row {i+2} forecast hour {j+1}"
                )
        return forecasts
    finally:
        workbook.close()


def load_inputs(workspace: Path) -> LoadedInputs:
    """Load and validate all 2025 source arrays; do not write any workbook."""
    root = Path(workspace).resolve()
    attachment_dir = root / "C题" / "附件"
    required = [
        attachment_dir / "附件2.xlsx",
        attachment_dir / "附件3.xlsx",
        attachment_dir / "附件4.xlsx",
        attachment_dir / "附件5" / "result4-2.xlsx",
        attachment_dir / "附件5" / "result4-3.xlsx",
    ]
    for path in required:
        if not path.is_file():
            raise FileNotFoundError(path)
    load_kwh, pv_kwh = _read_actuals(required[0])
    forecasts = _read_pv_forecasts(required[1])
    prices = _read_prices(required[2])
    days = [
        DayData(
            FIRST_DAY + timedelta(days=i),
            load_kwh[i], pv_kwh[i], prices[i], forecasts[i],
        )
        for i in range(DAY_COUNT)
    ]
    return LoadedInputs(days, required[3], required[4])
