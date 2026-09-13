"""Extract verified fourth-question data and build the requested editable Word tables."""
from pathlib import Path
from datetime import datetime
import csv
import hashlib
import json
import math

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DEST = Path(__file__).resolve().parents[1]
SOURCE = DEST.parent / 'output'
DATES = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']
CAP1 = '表1  微网在指定时间段的购电量及全天的购电量和购电费'
CAP2 = '表2  储能设备在指定时间段的充放电量及0:00和24:00的储电量'
CAP3 = '表3  微网在指定日期的紧急购电量'
NOTE = '表1购电量为最终正常购电量；全天购电费包含计划、调整及紧急购电费用。'


def clock(t):
    return f'{t // 6}:{t % 6 * 10:02d}'


def close(a, b):
    assert math.isclose(float(a), float(b), rel_tol=0, abs_tol=1e-6), (a, b)


def extract(strategy):
    paths = [SOURCE / f'checkpoint_{strategy}.jsonl', SOURCE / f'detail_{strategy}.csv']
    records = {r['date']: r for r in map(json.loads, paths[0].read_text(encoding='utf-8').splitlines()) if r['date'] in DATES}
    with paths[1].open(encoding='utf-8-sig', newline='') as stream:
        detail = {d: [] for d in DATES}
        for row in csv.DictReader(stream):
            if row['date'] in detail:
                detail[row['date']].append(row)
    days = []
    for date in DATES:
        r, rows = records[date], sorted(detail[date], key=lambda x: int(x['t']))
        assert [int(row['t']) for row in rows] == list(range(144))
        for t, row in enumerate(rows):
            assert row['interval'] == f'{clock(t)}-{clock(t+1)}'
            for key in ('initial_purchase', 'final_purchase', 'charge', 'discharge', 'emergency'):
                close(row[key+'_kwh'], r[key][t])
            close(row['soc_before_kwh'], r['soc_path'][t])
            close(row['soc_after_kwh'], r['soc_path'][t+1])
        for key in ('plan_cost', 'adjustment_cost', 'emergency_cost', 'total_cost'):
            close(sum(float(row[key+'_yuan']) for row in rows), r[key])
        close(r['soc_start'], r['soc_path'][0])
        close(r['soc_end'], r['soc_path'][-1])
        intervals = []
        start = None
        for t, amount in enumerate(r['emergency']+[0]):
            if amount > 1e-8 and start is None:
                start = t
            elif amount <= 1e-8 and start is not None:
                intervals.append({'start': start, 'end': t, 'time': f'{clock(start)}-{clock(t)}', 'kwh': sum(r['emergency'][start:t])})
                start = None
        close(sum(e['kwh'] for e in intervals), sum(r['emergency']))
        days.append({
            'date': date, 'date_label': f'{int(date[:4])}.{int(date[5:7])}.{int(date[8:])}',
            'selected': [{'time': f'{h}:00-{h}:10', 't': h*6, 'kwh': r['final_purchase'][h*6]} for h in (10,12,14,16,18,20)],
            'blocks': [{'time': f'{b*4}:00-{(b+1)*4}:00', 'charge': sum(r['charge'][24*b:24*(b+1)]), 'discharge': sum(r['discharge'][24*b:24*(b+1)])} for b in range(6)],
            'total_normal_kwh': sum(r['final_purchase']), 'total_cost_yuan': r['total_cost'],
            'soc_start': r['soc_start'], 'soc_end': r['soc_end'], 'emergency': intervals,
            'detail': [{k: (v if k in ('date','interval') else float(v)) for k,v in row.items()} for row in rows],
        })
    return {'strategy': strategy, 'table1_note': NOTE, 'captions': [CAP1,CAP2,CAP3], 'days': days,
            'sources': {str(p.relative_to(DEST.parent)).replace('\\','/'): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}}


def f4(value):
    return f'{0.0 if abs(value) < .00005 else value:.4f}'


def font(run, size=11, bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0,0,0)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key, value in [('ascii','Times New Roman'), ('hAnsi','Times New Roman'), ('eastAsia','黑体' if bold else '宋体')]:
        rf.set(qn('w:'+key), value)


def para(doc, text, size=11, bold=False, center=False, before=0, after=6, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = keep
    font(p.add_run(text), size, bold)
    return p


def cell_text(cell, text, size=11):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    font(p.add_run(str(text)), size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_table(doc, count, widths, height=23):
    table = doc.add_table(rows=count, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for row in table.rows:
        row.height = Pt(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in [('top',45),('bottom',45),('left',35),('right',35)]:
                el = OxmlElement('w:'+side)
                el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa'); margins.append(el)
            props.append(margins)
            cell_text(cell,'')
    return table


def build_docx(payload):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    for name in ('Normal','Title'):
        sty = doc.styles[name]
        sty.font.name = 'Times New Roman'
        sty.font.color.rgb = RGBColor(0,0,0)
        sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    strategy = payload['strategy']
    doc.core_properties.title = f'第四问策略{strategy}指定日期结果表'
    doc.core_properties.subject = '正常购电 储能充放电 紧急购电'
    doc.core_properties.author = ''
    for i, day in enumerate(payload['days']):
        if i in (0,2):
            title = para(doc, f'第四问策略{strategy}指定日期结果表', 16, True, True, after=8, keep=True, style='Title')
            title.paragraph_format.page_break_before = i == 2
            para(doc, '电量单位：kWh；费用单位：元。所有数值保留四位小数，汇总按未舍入值计算。', 9.5, after=5, keep=True)
            para(doc, NOTE, 9.5, after=9, keep=True)
        para(doc, day['date_label'], 13, True, True, before=14 if i%2 else 0, after=6, keep=True)
        para(doc, CAP1, 12, True, True, after=6, keep=True)
        t1 = make_table(doc, 4, [3]*6)
        for c, label in enumerate(['时间段','购电量']*3):
            cell_text(t1.cell(0,c), label)
        for j, item in enumerate(day['selected']):
            row, col = 1+j//3, 2*(j%3)
            cell_text(t1.cell(row,col),item['time'])
            cell_text(t1.cell(row,col+1),f4(item['kwh']))
        cell_text(t1.cell(3,0).merge(t1.cell(3,1)), '全天购电量')
        cell_text(t1.cell(3,2),f4(day['total_normal_kwh']))
        cell_text(t1.cell(3,3).merge(t1.cell(3,4)), '全天购电费')
        cell_text(t1.cell(3,5),f4(day['total_cost_yuan']))
        para(doc, CAP2, 12, True, True, before=11, after=6, keep=True)
        t2 = make_table(doc, 5, [3]*6)
        for c,label in enumerate(['时间段','充电量','放电量']*2):
            cell_text(t2.cell(0,c),label)
        for j,item in enumerate(day['blocks']):
            row,col = 1+j//2,3*(j%2)
            cell_text(t2.cell(row,col),item['time'])
            cell_text(t2.cell(row,col+1),f4(item['charge']))
            cell_text(t2.cell(row,col+2),f4(item['discharge']))
        cell_text(t2.cell(4,0).merge(t2.cell(4,1)), '0:00 储电量')
        cell_text(t2.cell(4,2),f4(day['soc_start']))
        cell_text(t2.cell(4,3).merge(t2.cell(4,4)), '24:00 储电量')
        cell_text(t2.cell(4,5),f4(day['soc_end']))
    title = para(doc, f'第四问策略{strategy}', 13, True, True, after=10, keep=True)
    title.paragraph_format.page_break_before = True
    para(doc, CAP3, 15, True, True, after=10, keep=True, style='Title')
    para(doc, '电量单位：kWh', 10, after=6, keep=True)
    n = max(len(d['emergency']) for d in payload['days'])
    t3 = make_table(doc, n+2, [2.35,2.15]*4, height=24)
    for k,day in enumerate(payload['days']):
        cell_text(t3.cell(0,2*k).merge(t3.cell(0,2*k+1)),day['date_label'],12)
        cell_text(t3.cell(1,2*k),'时间段')
        cell_text(t3.cell(1,2*k+1),'购电量')
        for j,item in enumerate(day['emergency']):
            cell_text(t3.cell(j+2,2*k),item['time'],10.5)
            cell_text(t3.cell(j+2,2*k+1),f4(item['kwh']),10.5)
    for row in t3.rows[:2]:
        row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    para(doc, '注：连续发生的10分钟紧急购电时段合并列示，购电量为该连续时间段内的合计。空白表示该日期无更多记录。',10,before=9,after=5)
    # Explicit black cell borders override conditional borders in the bundled template.
    for tree in (doc.styles.element,doc.element):
        for border in list(tree.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    for table in doc.tables:
        for tc in table._tbl.xpath('.//w:tc'):
            props = tc.get_or_add_tcPr()
            for old in list(props.findall(qn('w:tcBorders'))):
                props.remove(old)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                b=OxmlElement('w:'+side)
                b.set(qn('w:val'),'single');b.set(qn('w:sz'),'8');b.set(qn('w:color'),'000000')
                borders.append(b)
            props.append(borders)
    path=DEST/f'第四问_{strategy}_指定日期表1表2表3.docx'
    doc.save(path)
    assert len(Document(path).tables)==9
    return path


if __name__ == '__main__':
    (DEST/'data').mkdir(parents=True,exist_ok=True)
    for strategy in ('4-2','4-3'):
        payload=extract(strategy)
        (DEST/'data'/f'tables_{strategy}.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
        path=build_docx(payload)
        print(json.dumps({'path':str(path),'tables':9,'emergency_counts':[len(d['emergency']) for d in payload['days']]},ensure_ascii=False))
